from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


BODY_STYLE = "论文正文"
H2_STYLE = "heading 2"
H1_STYLE = "heading 1"


TOP_UP = {
    "系统总体架构设计": [
        "因此，本文的总体架构并不是简单罗列技术名词，而是围绕工业异常检测的完整任务链路组织。用户提出问题，系统接收任务，图结构安排执行顺序，Agent 分工完成检测和解释，工具层提供算法与知识能力，状态层保存过程证据，前端再把这些结果展示出来。每一层都有明确输入和输出，这也是后续详细设计能够展开的基础。",
    ],
    "视觉检测与PatchCore实现": [
        "在实际调试中，视觉模块也是最容易暴露问题的部分。模型配置、图片编码、类别选择、阈值设置、热力图保存路径都会影响最终结果。本文将这些信息尽量写入 metadata，就是为了让问题出现时能够回溯。例如 selected_backend 可以说明使用了哪种后端，threshold 可以说明当前阈值，heatmap_path 可以说明可视化文件是否生成。",
    ],
    "RAG知识库与Few-shot实现": [
        "RAG 与 few-shot 的结合也让系统具备了一定的可解释调试能力。若模型判断异常但 few-shot 正常参考更相似，开发者就需要重新检查 prompt 或阈值；若检索到的异常样本与当前图像明显不相关，则需要优化数据描述或检索条件。也就是说，RAG 不只是增强回答，也为后续误报分析提供了线索。",
    ],
    "多Agent流程与追问测试": [
        "测试追问链路时，还可以观察对话历史是否被压缩和保留。用户连续多轮提问后，系统不应无限堆叠完整历史，也不应丢失首轮异常结论。通过 conversation_summary 和最近对话共同参与回答，系统能在成本和上下文完整性之间取得折中。这一点对于长任务排障尤其重要。",
    ],
}


def paragraph_text(paragraph) -> str:
    return paragraph.text.strip()


def style_name(paragraph) -> str:
    try:
        return paragraph.style.name.lower()
    except Exception:
        return ""


def find_heading(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph_text(paragraph) == text and style_name(paragraph) == H2_STYLE:
            return paragraph
    raise RuntimeError(text)


def find_next_heading(doc: Document, anchor):
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p._element is anchor._element)
    for paragraph in paragraphs[start + 1 :]:
        if style_name(paragraph) in {H1_STYLE, H2_STYLE} or paragraph_text(paragraph) == "参考文献":
            return paragraph
    raise RuntimeError("next heading not found")


def main() -> int:
    if len(sys.argv) != 2:
        return 2
    path = Path(sys.argv[1])
    doc = Document(str(path))
    for heading, paragraphs in TOP_UP.items():
        anchor = find_heading(doc, heading)
        stop = find_next_heading(doc, anchor)
        for text in paragraphs:
            para = stop.insert_paragraph_before(text)
            try:
                para.style = BODY_STYLE
            except Exception:
                pass
            for run in para.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
    doc.save(str(path))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
