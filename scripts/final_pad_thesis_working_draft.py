from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


BODY_STYLE = "论文正文"
H1_STYLE = "heading 1"
H2_STYLE = "heading 2"


PADDING: dict[str, list[str]] = {
    "论文组织结构": [
        "这样的章节安排也对应了系统开发的自然顺序。先从问题背景和相关技术说明为什么要做，再从需求和总体架构说明系统要长成什么样，随后在详细设计中说明各模块如何落到代码，最后通过测试和实验分析验证系统是否真的能够运行。把需求分析与总体设计放在同一章后，论文结构更紧凑，也能减少原稿中第二、三、四章之间的重复。",
    ],
    "PatchCore异常定位方法": [
        "在论文表述中，还可以把 PatchCore 与通用视觉模型作一个工程层面的对比。PatchCore 更像一个“类别内异常距离计算器”，它不会理解用户的自然语言问题，却能在同类别正常样本充分时给出较直观的热区；通用视觉模型更像一个“图文解释器”，它能回答用户问题，却可能在细粒度缺陷上不稳定。本文系统把两者接入同一工具层，目的就是让用户根据任务场景选择更合适的后端。",
    ],
    "RAG与少样本上下文": [
        "RAG 知识库的构建质量会直接影响 few-shot 上下文。若样本描述过于简单，检索结果可能只按类别匹配，无法体现缺陷差异；若描述中包含缺陷类型、位置、严重程度和图像路径，检索结果就更容易被 KnowledgeAgent 使用。因此，本文后续实验除了看是否能检索到样本，还要观察返回样本是否具有足够的语义信息。",
    ],
    "LangGraph与多Agent协作": [
        "与传统线性流程相比，图编排还有一个优点：它能把等待用户、重试和报告生成这类非主线操作纳入同一套状态流转。工业检测并不总是从输入到输出一次完成，系统经常需要等用户确认区域、重新选择阈值或补充问题。把这些分支写进图结构后，系统行为更容易解释，也更方便在前端以时间线方式展示。",
    ],
    "功能需求分析": [
        "对于前端来说，功能需求还包括状态反馈。用户点击“开始检测”后，如果后端长时间没有响应，用户并不知道任务是在上传、检测、检索还是报告生成阶段。因此，系统通过流式接口把事件持续推送到前端，让用户看到 Agent 正在执行的步骤。这个需求来自长任务排障场景，也是多 Agent 系统比普通接口更需要重视的地方。",
    ],
    "系统总体架构设计": [
        "状态记忆层虽然放在架构图下方，但它并不是被动存储。Supervisor 规划时需要读取历史问题和已有检测结果，KnowledgeAgent 需要读取 VisionAgent 的异常信息，AnswerNode 需要读取 RAG 上下文和记忆摘要，ReportAgent 需要读取整个任务的结构化结果。可以说，状态记忆层把原本松散的模块连接成了一个连续任务。",
    ],
    "主运行流程设计": [
        "从用户体验角度看，主流程还要避免两个极端：一是每个问题都调用所有 Agent，导致响应慢、成本高；二是只调用视觉模型，导致后续解释和追问能力不足。Supervisor 的作用就是在这两者之间做选择。它根据用户意图和已有状态决定最小必要步骤，使系统既能保持灵活，也能避免不必要的工具调用。",
    ],
    "API接口与任务模型实现": [
        "接口层还需要处理跨域和静态资源。当前后端配置了 CORS，并挂载了 /web、/data/uploads 和 /data/heatmaps 等静态路径。这样，本地前端或通过 SSH 隧道访问服务器时，可以直接请求后端接口，也可以加载上传图片和热力图结果。对于部署到 AutoDL 后在本地浏览器调试的场景，这一点非常关键。",
    ],
    "图编排与Supervisor实现": [
        "执行事件 execution_events 是系统可观测性的关键字段。它记录每个步骤的开始、结束、状态和可能的错误。前端时间线并不是凭空生成的，而是读取这些事件后进行渲染。因此，论文中可以把时间线展示和后端状态字段联系起来说明，体现前后端不是割裂实现，而是围绕同一运行状态协作。",
    ],
    "视觉检测与PatchCore实现": [
        "在异常区域解释方面，系统没有只保存热力图，而是进一步根据区域面积、长宽比例和热度强度推导 size、shape、texture、severity_hint 等字段。这些字段虽然是规则生成的，但它们让报告中的文字更具体。例如，同样是异常热区，系统可以区分小型紧凑块状区域和细长线性响应，从而让缺陷描述更接近人工复核语言。",
    ],
    "RAG知识库与Few-shot实现": [
        "RAG 运行失败时，系统不应中断基础检测。当前设计中，如果知识库尚未构建、向量库为空或检索失败，系统仍可以继续执行视觉检测，并在 metadata 或回答中保留相应提示。这种降级策略很重要，因为工业检测首先要保证基本可用，然后再逐步增强解释能力。",
    ],
    "状态记忆与追问续接实现": [
        "上下文恢复还涉及 pending clarification。若 ClarificationAgent 生成了澄清问题，系统需要保存 pending_question 和 pending_clarification，等待用户补充后继续流程。此前代码中曾出现 clarification 为 None 时直接取 get 导致异常的问题，修复后系统能更稳地处理空上下文。这类细节虽然不显眼，但能体现状态恢复模块的工程必要性。",
    ],
    "前端交互与报告生成实现": [
        "前端热力图展示的价值在于把算法结果变成可检查证据。用户可以同时看到原图、热力图、叠加图和异常描述，再判断系统定位是否合理。如果只有文字回答，用户很难知道异常位置是否对应真实缺陷；如果只有热力图，没有文字描述，用户又需要自己解释热区含义。因此，前端把图像证据和文本解释放在一起展示。",
    ],
    "测试环境与数据来源": [
        "AutoDL 服务器实验还需要注意 CUDA 与 PyTorch 版本兼容性。前期训练时出现过 RTX 5090 与当前 PyTorch 安装不兼容的提示，最后通过 CPU 或更合适的依赖组合完成训练。这说明实验环境不是可省略信息，尤其是涉及 PatchCore 特征提取和 GPU 加速时，硬件、驱动和框架版本都会影响复现。",
    ],
    "接口与功能测试": [
        "测试用例还应覆盖前端编码和中文显示。项目曾出现前端乱码、USER.md 编码清洗、Linux 终端乱码等问题。对于中文论文系统来说，如果前端按钮、时间线、Markdown 渲染或错误提示出现乱码，会直接影响演示效果。因此，编码一致性也可以作为前端功能测试的一部分。",
    ],
    "PatchCore异常定位实验": [
        "热力图实验最终需要落到图片证据。论文中建议选择一张 good 样本和一张 broken_small 样本，分别展示 overlay 图，并在文字中说明阈值变化带来的结果差异。这样读者不需要理解全部代码，也能看出 PatchCore 工具的作用：它为异常定位提供了可视化依据，而不仅是返回一个异常分数。",
    ],
    "RAG检索与Few-shot实验": [
        "RAG 查询还可以设计“同类别”和“跨类别”两个场景。同类别查询用于验证知识库能否召回更接近的正常/异常样本；跨类别查询用于观察检索结果是否会偏离当前任务。若跨类别结果较多，说明查询文本或 metadata 过滤还需要优化。这样的分析能让第五章更像实验，而不是只写功能已经实现。",
    ],
    "多Agent流程与追问测试": [
        "多 Agent 时间线截图建议放在本节。截图中应能看到从 Supervisor 计划到 VisionAgent、KnowledgeAgent 或 ReportAgent 执行的过程。若某个节点失败，也应展示错误信息如何传递到前端。这样的截图比单纯最终回答更能证明系统确实采用了多 Agent 编排，而不是只在论文中描述了多个角色名称。",
    ],
    "工作总结": [
        "论文最终总结时，可以强调本文系统的定位是“面向工业异常诊断流程的应用型系统”。它关注的是如何让异常检测从模型输出走向可交互诊断，而不是单独提出一个新算法。这样的定位能够把前端、后端、RAG、PatchCore、多 Agent 和报告生成都纳入同一个贡献框架。",
    ],
    "系统不足": [
        "还需要指出的是，当前系统对真实工业现场的适配仍有限。实际产线会有相机标定、光源稳定性、节拍约束、缺陷复核标准和人工闭环流程，这些并未在当前项目中完整实现。因此，论文应把系统描述为实验性原型和毕业设计实现，而不是已经可直接上线的工业质检平台。",
    ],
    "后续优化方向": [
        "后续还可以把前端操作反馈纳入在线学习闭环。例如，用户确认某个异常为误报后，系统可以把该样本作为正常参考写入知识库；用户确认某个热区为真实缺陷后，系统可以补充缺陷描述和类别标签。这样，RAG 知识库就能随着使用逐步积累经验，为后续 few-shot 和知识分析提供更可靠的样本。",
    ],
}


def paragraph_text(paragraph) -> str:
    return paragraph.text.strip()


def style_name(paragraph) -> str:
    try:
        return paragraph.style.name.lower()
    except Exception:
        return ""


def find_body_heading(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph_text(paragraph) == text and style_name(paragraph) in {H1_STYLE, H2_STYLE}:
            return paragraph
    raise RuntimeError(f"heading not found: {text}")


def find_next_body_heading(doc: Document, heading_para):
    paragraphs = doc.paragraphs
    start = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph._element is heading_para._element
    )
    for paragraph in paragraphs[start + 1 :]:
        if style_name(paragraph) in {H1_STYLE, H2_STYLE} or paragraph_text(paragraph) == "参考文献":
            return paragraph
    raise RuntimeError(f"next heading not found after: {paragraph_text(heading_para)}")


def add_para_before(ref_para, text: str):
    para = ref_para.insert_paragraph_before(text)
    try:
        para.style = BODY_STYLE
    except Exception:
        pass
    for run in para.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: final_pad_thesis_working_draft.py <docx_path>")
        return 2
    path = Path(sys.argv[1])
    doc = Document(str(path))
    for heading, paragraphs in PADDING.items():
        anchor = find_body_heading(doc, heading)
        stop = find_next_body_heading(doc, anchor)
        for text in paragraphs:
            add_para_before(stop, text)
    doc.save(str(path))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
