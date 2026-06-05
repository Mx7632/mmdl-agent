from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor


BODY_STYLE = "论文正文"
H1_STYLE = "heading 1"
H2_STYLE = "heading 2"
H3_STYLE = "heading 3"
TOC1_STYLE = "toc 1"
TOC2_STYLE = "toc 2"
TOC3_STYLE = "toc 3"


CN_ABSTRACT = [
    "随着工业制造质量控制要求不断提高，工业异常检测任务已经不只是判断一张图像是否存在缺陷。实际检测人员往往还需要知道异常大致出现在什么位置、可能属于哪类缺陷、是否会影响产品功能，以及后续应当怎样复核和处置。传统视觉检测方法通常更擅长给出异常分数或热力图，而通用多模态大模型虽然具备较好的图像理解和语言表达能力，却容易受到正常纹理、反光、阴影和背景干扰影响。因此，如何把视觉检测、知识检索、人工追问和报告生成放在同一条可追踪流程中，是本文关注的主要问题。",
    "针对上述问题，本文基于当前项目实现了一套面向工业异常检测的多 Agent 与检索增强生成系统。系统后端采用 FastAPI 提供检测、流式检测、RAG 建库、知识查询和报告生成等接口；任务执行过程由 LangGraph 组织为数据加载、Supervisor 规划、专家执行、结果合并、自我反思、等待用户和回答输出等节点。视觉检测侧支持通用视觉模型、专业 HTTP 检测服务和 PatchCore 本地检测工具；其中 PatchCore 通过正常样本构建特征库，并在推理时生成热力图、mask、overlay 和异常区域描述。",
    "在知识增强方面，系统通过 RAG 检索相似案例，并由 FewShotSelector 从候选结果中筛选正常样本和异常样本，构造少样本参考上下文，再注入 VisionAgent 或 KnowledgeAgent。KnowledgeAgent 进一步生成缺陷成因、风险提示、建议动作和产品分析等结构化内容。系统还引入 MMAD 七类任务思想，将异常判别、缺陷分类、缺陷定位、缺陷描述、缺陷分析、产品分类和产品分析统一到诊断结果中。测试章节以功能验证、接口验证、PatchCore 示例、RAG 检索和多 Agent 调度为主，避免对尚未系统评测的性能指标作过度结论。",
]


EN_ABSTRACT = [
    "With the increasing demand for quality control in industrial manufacturing, industrial anomaly detection is no longer limited to a simple normal-or-defective decision. In practical inspection scenarios, users also need coarse localization, defect type information, possible causes, risk hints, follow-up suggestions, and a traceable report. Traditional visual anomaly detection methods are good at anomaly scores or heatmaps, while general multimodal models provide stronger visual-language explanations but may be affected by reflections, shadows, normal textures, and background noise. This thesis therefore focuses on integrating visual detection, knowledge retrieval, human clarification, and report generation into one traceable workflow.",
    "Based on the current project, this thesis designs and implements a multi-agent industrial anomaly detection system with retrieval-augmented generation. The backend uses FastAPI to provide detection, streaming detection, RAG building, knowledge query, and report generation APIs. The workflow is organized by LangGraph, including data loading, supervisor planning, expert execution, result merging, self-reflection, user-waiting, and answer generation. The visual side supports a general vision model, an external professional detector, and a local PatchCore tool. PatchCore builds a memory bank from normal samples and produces heatmaps, masks, overlays, and localized anomaly descriptions during inference.",
    "For knowledge enhancement, the system retrieves similar cases through RAG. FewShotSelector selects normal and abnormal examples from retrieved candidates, and FewShotPromptBuilder constructs a compact few-shot context for VisionAgent or KnowledgeAgent. KnowledgeAgent further generates structured defect causes, risk notes, repair suggestions, and object analysis. The system also adopts the seven-task idea of MMAD to organize anomaly discrimination, defect classification, defect localization, defect description, defect analysis, object classification, and object analysis. The evaluation chapter is designed around functional verification, API testing, PatchCore examples, RAG retrieval, and multi-agent routing, without fabricating unsupported performance metrics.",
]


TOC_ENTRIES = [
    (1, "第一章 绪论"),
    (2, "1.1 研究背景与意义"),
    (2, "1.2 国内外研究现状"),
    (2, "1.3 现有问题分析"),
    (2, "1.4 本文主要工作"),
    (2, "1.5 论文组织结构"),
    (1, "第二章 相关技术"),
    (2, "2.1 工业异常检测技术"),
    (2, "2.2 PatchCore异常定位方法"),
    (2, "2.3 多模态大模型与MMAD任务"),
    (2, "2.4 RAG与少样本上下文"),
    (2, "2.5 LangGraph与多Agent协作"),
    (2, "2.6 本章小结"),
    (1, "第三章 需求分析与系统总体设计"),
    (2, "3.1 应用场景与使用角色"),
    (2, "3.2 功能需求分析"),
    (2, "3.3 非功能需求分析"),
    (2, "3.4 系统总体架构设计"),
    (2, "3.5 主运行流程设计"),
    (2, "3.6 本章小结"),
    (1, "第四章 系统详细设计与实现"),
    (2, "4.1 API接口与任务模型实现"),
    (2, "4.2 图编排与Supervisor实现"),
    (2, "4.3 视觉检测与PatchCore实现"),
    (2, "4.4 RAG知识库与Few-shot实现"),
    (2, "4.5 状态记忆与追问续接实现"),
    (2, "4.6 前端交互与报告生成实现"),
    (1, "第五章 系统测试与实验分析"),
    (2, "5.1 测试环境与数据来源"),
    (2, "5.2 接口与功能测试"),
    (2, "5.3 PatchCore异常定位实验"),
    (2, "5.4 RAG检索与Few-shot实验"),
    (2, "5.5 多Agent流程与追问测试"),
    (2, "5.6 本章小结"),
    (1, "第六章 总结与展望"),
    (2, "6.1 工作总结"),
    (2, "6.2 系统不足"),
    (2, "6.3 后续优化方向"),
    (2, "6.4 本章小结"),
    (1, "参考文献"),
    (1, "致　　谢"),
]


BODY_BLOCKS = [
    ("h1", "绪论"),
    ("h2", "研究背景与意义"),
    ("p", "工业异常检测是制造质量控制中的重要环节。对瓶口、金属表面、螺丝、线缆、药片等产品来说，缺陷往往只占图像中的很小一部分，却可能影响产品装配、密封、外观或后续使用安全。如果只依赖人工目检，不仅效率容易受到生产节拍影响，也难以保证长时间检测的一致性。因此，利用计算机视觉和智能系统辅助检测，已经成为工业质检中很自然的需求。"),
    ("p", "不过，真实的异常检测并不是一句“有异常”或“无异常”就能结束。检测人员通常会继续追问：异常在哪里？像划痕、破损还是污染？这个缺陷可能来自加工、运输还是成像环境？如果要复核，应当看哪一块区域？这也是本文把异常检测系统设计成多 Agent 工作流的原因。系统不仅需要看图，还需要把检测证据、知识解释、人工澄清和报告输出连成一条可追踪链路。"),
    ("h2", "国内外研究现状"),
    ("p", "在工业异常检测领域，MVTec AD 等数据集推动了无监督异常检测方法的发展。PaDiM、PatchCore、STFPM、DRAEM、EfficientAD 等方法从不同角度利用正常样本建模或特征差异进行异常检测与定位。其中 PatchCore 通过保存正常样本的局部特征，并在推理时计算测试图像局部特征与正常特征库之间的距离，能够输出异常分数和热力图，是工业异常定位中较有代表性的一类方法。"),
    ("p", "多模态大模型的发展让异常检测多了一种新的表达方式。相比传统模型只输出分数或掩码，多模态模型可以用自然语言描述缺陷位置、形态和可能影响。MMAD 论文进一步把工业异常理解拆成异常判别、缺陷分类、缺陷定位、缺陷描述、缺陷分析、产品分类和产品分析七类任务，这对本文组织系统输出有直接启发。与此同时，RAG 和 Agent 技术也为系统引入外部案例知识、工具调用和多轮追问提供了工程基础。"),
    ("h2", "现有问题分析"),
    ("p", "现有方案主要存在三类问题。第一，传统异常检测模型通常强调检测结果和定位效果，但对异常原因、风险和处理建议支持不足。第二，通用视觉大模型虽然能描述图像，但在工业场景中容易把反光、阴影、正常纹理和边缘误判为缺陷。第三，很多系统仍停留在单次模型调用层面，缺少任务状态、追问续接、知识检索和报告沉淀机制。"),
    ("p", "这些问题说明，工业异常检测系统不能只追求一个模型输出，而应当把不同能力组合起来。PatchCore 更适合固定类别的本地异常定位，通用视觉模型更适合开放描述，RAG 能补充相似案例和专业知识，多 Agent 编排则负责决定什么时候调用哪个能力。本文的系统设计正是围绕这种组合式思路展开。"),
    ("h2", "本文主要工作"),
    ("p", "本文完成的主要工作包括以下几个方面。第一，设计了以 Supervisor 为核心的多 Agent 工业异常检测流程，把视觉检测、知识分析、人工澄清、回答生成和报告生成拆成不同职责。第二，实现了基于 FastAPI 和 LangGraph 的任务运行链路，使前端上传图像后可以通过流式接口观察执行过程。第三，接入 Qwen 通用视觉检测、专业 HTTP 视觉服务和 PatchCore 本地检测工具，并为 PatchCore 增加训练、推理、类别检查和热力图输出能力。"),
    ("p", "第四，构建了 RAG 知识库与 few-shot 上下文链路，支持从数据集中检索相似正常/异常案例，再将筛选后的参考样本注入视觉或知识分析流程。第五，结合 MMAD 七任务思想，将检测结果扩展为异常判别、缺陷定位、缺陷描述、缺陷分析和产品分析等结构化内容。第六，前端实现了检测、追问、RAG 建库、热力图预览和多 Agent 时间线展示，为系统调试和论文实验提供了可观察入口。"),
    ("h2", "论文组织结构"),
    ("p", "全文共分为六章。第一章介绍研究背景、研究现状、问题分析和主要工作。第二章说明工业异常检测、PatchCore、多模态大模型、RAG 和多 Agent 编排等相关技术。第三章把需求分析与系统总体设计放在一起，先说明系统要解决什么问题，再说明系统整体如何分层。第四章展开系统详细设计与实现，重点对应代码中的 API、图编排、视觉检测、RAG、状态记忆和前端模块。第五章给出系统测试与实验分析方案，并基于真实运行链路进行功能验证。第六章总结本文工作，并说明系统仍存在的不足和后续优化方向。"),
    ("h1", "相关技术"),
    ("h2", "工业异常检测技术"),
    ("p", "工业异常检测通常面向“正常样本较多、异常样本较少”的场景。与一般图像分类不同，工业异常往往种类多、形态细小，而且不同产品类别之间差异很大。因此，许多方法会先学习正常样本的特征分布，再把偏离正常分布的区域视为异常候选。这样做的好处是能够减少对大量异常样本的依赖，更符合工业数据采集的现实情况。"),
    ("p", "在本文系统中，工业异常检测既包含图像级判断，也包含区域级定位和文字解释。图像级判断回答“是否异常”，区域级定位回答“异常在哪里”，文字解释则回答“异常可能意味着什么”。这三者并不是互相替代的关系，而是共同构成一次完整诊断。"),
    ("h2", "PatchCore异常定位方法"),
    ("p", "PatchCore 是本文本地热力图工具的核心思路。它不直接训练一个分类器区分正常和异常，而是利用预训练卷积网络提取正常样本的 patch 特征，并将这些特征保存为 memory bank。推理时，系统对输入图像提取同样的 patch 特征，再计算这些特征与正常 memory bank 的距离。距离越大，说明该局部区域越不像训练集中见过的正常外观。"),
    ("p", "当前项目中的 PatchCore 实现包含训练脚本和推理脚本。训练脚本读取 MVTec 数据集中某一类别的 train/good 图像，生成 memory_bank.pt 和 metadata.json；推理脚本读取待测图像、类别和阈值，输出异常区域、热力图、叠加图、mask 和相关元数据。由于 memory bank 与产品类别绑定，PatchCore 更适合已经训练过的固定类别。对于没有训练过的类别，系统会通过类别检查逻辑提示用户先训练该类别或切换到通用视觉后端。"),
    ("h2", "多模态大模型与MMAD任务"),
    ("p", "多模态大模型能够同时处理图像和文本，因此适合承担异常描述、缺陷解释和交互问答等任务。本文并不把通用视觉模型当作唯一检测依据，而是把它作为视觉后端之一。系统要求模型返回结构化 JSON，包括异常列表、位置、置信分数和观测说明。这样可以让后续流程继续解析，而不是停留在一段不可控的自然语言回答上。"),
    ("p", "MMAD 对本文的价值主要体现在任务组织上。它将工业异常理解拆成七个子任务，提醒系统不能只判断异常有无，还需要考虑缺陷类别、位置、描述、成因和产品影响。本文在报告与结构化诊断部分借鉴了这一思想，把视觉检测结果和知识分析结果统一归纳到多个任务字段中，便于后续展示和报告生成。"),
    ("h2", "RAG与少样本上下文"),
    ("p", "RAG 的基本思想是先从外部知识库检索相关内容，再把检索结果注入模型上下文。对于工业异常检测来说，知识库可以保存数据集样本、缺陷描述、相似案例、正常外观说明和维修建议。这样，模型在回答时不只依赖自身参数知识，还能参考当前项目中已经整理好的案例。"),
    ("p", "本文系统中的 few-shot 不是重新训练模型，而是提示级的上下文增强。系统从 RAG 检索结果中选择少量正常样本和异常样本，构造紧凑的参考文本。VisionAgent 可以据此判断当前图像更接近正常参考还是异常参考；KnowledgeAgent 则可以使用这些案例生成更具体的缺陷分析。需要注意的是，当前实现主要是基于检索结果和元数据的样本筛选，并不是复杂的监督学习评分模型。"),
    ("h2", "LangGraph与多Agent协作"),
    ("p", "LangGraph 用于描述任务执行的节点和条件路由。相比把所有逻辑写进一个函数，图结构更适合多步骤、多分支和可恢复的任务。本文系统将任务拆成 load_data、supervisor_plan、supervisor_execute、supervisor_merge、self_reflect、wait_user、answer 和 report 等节点，并根据执行状态决定下一步走向。"),
    ("p", "多 Agent 协作的核心不是让多个模型随意对话，而是让不同职责有清晰边界。Supervisor 负责计划和路由，VisionAgent 负责视觉检测，KnowledgeAgent 负责知识分析，ClarificationAgent 负责不确定时的追问，ReportAgent 负责报告组织，AnswerNode 负责最终回答。这种拆分让系统更容易调试，也更容易在前端时间线中展示每一步状态。"),
    ("h2", "本章小结"),
    ("p", "本章介绍了系统涉及的主要技术基础。工业异常检测提供了问题背景，PatchCore 提供了固定类别热力图能力，多模态大模型提供了图文理解和描述能力，RAG 与 few-shot 上下文提供了案例参考，多 Agent 与 LangGraph 则提供了任务编排方式。这些技术共同支撑了后续系统设计。"),
    ("h1", "需求分析与系统总体设计"),
    ("h2", "应用场景与使用角色"),
    ("p", "本文系统面向工业图像异常诊断场景。典型使用过程是：用户上传待检测图像，选择视觉后端或产品类别，系统执行检测并返回异常列表、位置、热力图或解释文本；如果用户继续询问原因、风险或处理建议，系统需要基于上一轮上下文继续回答；如果检测结果不确定，系统可以提出澄清问题；如果用户需要归档，系统可以生成结构化报告。"),
    ("p", "系统中的使用角色不是严格的账号权限角色，而是论文需求分析中的使用场景划分。普通检测人员关注检测结果和异常位置，工程技术人员关注缺陷原因和处理建议，系统维护人员关注模型、知识库、PatchCore 类别和运行状态，论文研究人员关注多 Agent 调度过程、RAG 证据和实验记录。"),
    ("h2", "功能需求分析"),
    ("p", "系统需要支持图像异常检测、异常定位、知识检索、追问续接、人工澄清、报告生成和前端可视化等功能。图像异常检测要求系统能够接收上传图像并输出结构化异常候选；异常定位要求在具备条件时返回 bbox、粗略方位或热力图；知识检索要求系统能够从 RAG 知识库中召回相似案例；追问续接要求系统保留上一轮任务上下文；报告生成要求系统把检测结果和知识分析整理成较完整的诊断文本。"),
    ("table", "表3-1 系统功能需求表", ["需求项", "说明", "项目证据"], [
        ["图像检测", "接收上传图像并返回异常候选、描述和元数据", "DetectionTask、/v1/detect、/v1/stream"],
        ["热力图定位", "对已训练 PatchCore 类别生成 heatmap、mask 和 overlay", "patchcore_detection.py、patchcore_eval.py"],
        ["RAG检索", "构建知识库并查询相似案例", "/v1/rag/build、/v1/rag/query"],
        ["追问续接", "基于 task_id 恢复上下文并继续回答", "/v1/chat、/v1/continue、DetectionState"],
        ["报告生成", "复用检测和知识结果生成诊断报告", "/v1/generate_report、ReportAgent"],
    ]),
    ("h2", "非功能需求分析"),
    ("p", "从论文系统角度看，非功能需求主要包括可解释性、可扩展性、可维护性和可观测性。可解释性要求系统不仅输出结果，还能展示异常位置、相似案例、知识来源和 Agent 执行过程。可扩展性要求视觉后端、RAG 知识库和 Agent 节点能够逐步替换或增加。可维护性要求数据模型、工具调用和状态结构分层清楚。可观测性要求前端和日志能够看到任务进度、节点状态、trace_id 和错误信息。"),
    ("p", "本文不将系统描述为已经具备大规模生产部署能力。当前实现更接近一个面向毕业设计和实验验证的工程原型，它重在把多 Agent、RAG、PatchCore 和前端可视化连成闭环，而不是证明系统在高并发或复杂工业现场中已经完成长期验证。"),
    ("h2", "系统总体架构设计"),
    ("p", "系统采用分层架构，整体包括用户交互层、API 与任务运行层、图编排层、多 Agent 协作层、工具与知识层、状态记忆层。用户交互层由前端页面承担，负责图像上传、后端选择、RAG 建库、结果展示和多 Agent 时间线展示。API 与任务运行层由 FastAPI 实现，负责接收请求、解析表单、保存上传数据、返回普通响应或 SSE 流式响应。"),
    ("p", "图编排层由 LangGraph 实现，控制任务从加载数据到回答输出的流转。多 Agent 协作层负责把视觉、知识、澄清和报告任务分配给不同专家节点。工具与知识层包含 Qwen 视觉检测、专业 HTTP 检测服务、PatchCore、本地 RAG、Chroma 向量库和 MMAD 分析。状态记忆层保存任务状态、对话历史、共享上下文、执行事件和报告请求，使追问和恢复成为可能。"),
    ("table", "表3-2 系统分层职责表", ["层次", "主要职责", "对应实现"], [
        ["用户交互层", "上传图像、选择后端、查看结果和时间线", "web/index.html"],
        ["API与任务运行层", "提供检测、流式、RAG、报告等接口", "app/api/main.py"],
        ["图编排层", "定义节点和条件路由", "app/core/graph.py"],
        ["多Agent协作层", "规划、执行、合并和反思", "app/agents、app/orchestration"],
        ["工具与知识层", "视觉检测、PatchCore、RAG、MMAD分析", "app/tools、app/rag、app/analysis"],
        ["状态记忆层", "任务状态、对话历史和共享上下文", "app/memory、context_store.py"],
    ]),
    ("h2", "主运行流程设计"),
    ("p", "系统主流程从用户请求开始。前端通过 /v1/stream 或 /v1/detect 发送任务，后端构造 DetectionTask，并将图像以 base64 形式放入参数中。LangGraph 首先执行 load_data 节点，然后进入 supervisor_plan。Supervisor 根据用户问题、输入类型、已有上下文和当前阶段生成执行计划。若存在计划步骤，流程进入 supervisor_execute；如果不需要专家执行，则直接进入合并阶段。"),
    ("p", "专家执行完成后，supervisor_merge 会将工具结果和 Agent 输出写入共享上下文。self_reflect 节点检查是否需要重试、澄清或直接回答。如果需要用户补充信息，流程进入 wait_user；如果可以回答，则进入 answer。若用户请求报告，answer 后还会进入 report 节点。这样的流程让系统能够处理首轮检测、追问、失败重试、人工澄清和报告生成等不同情况。"),
    ("h2", "本章小结"),
    ("p", "本章将需求分析与总体设计放在一起。这样安排的好处是逻辑更连贯：先说明系统面对哪些场景、需要哪些功能，再说明系统为什么采用分层架构和多 Agent 编排。后续第四章将进一步进入代码层面，说明这些设计如何在项目中实现。"),
    ("h1", "系统详细设计与实现"),
    ("h2", "API接口与任务模型实现"),
    ("p", "系统后端入口位于 app/api/main.py。FastAPI 应用提供健康检查、图像检测、流式检测、继续任务、追问聊天、生成报告、RAG 建库、RAG 查询和 PatchCore 类别查询等接口。对于图像检测，请求通常以 multipart/form-data 方式上传，后端读取 image 字段并编码为 base64，再将其放入 DetectionTask.parameters。"),
    ("p", "任务和结果模型定义在 app/schemas/detection.py。DetectionTask 保存 task_id、asset_id、时间范围、数据来源、输入类型、问题和参数；DetectionResult 保存状态、回答、异常列表、摘要、解释和元数据。RAG 相关请求与响应也在同一文件中定义，包括 RagBuildRequest、RagBuildStatusResponse、RagQueryRequest 和 RagQueryResponse 等。统一的数据模型减少了接口层和 Agent 层之间的格式转换成本。"),
    ("h2", "图编排与Supervisor实现"),
    ("p", "图编排的核心文件是 app/core/graph.py。系统将任务执行拆分为 load_data、supervisor_plan、supervisor_execute、supervisor_merge、self_reflect、wait_user、answer 和 report 节点，并通过条件边控制任务流向。例如，supervisor_plan 后如果存在计划步骤，就进入 supervisor_execute；否则进入 supervisor_merge。self_reflect 会根据反思结果决定重试、继续规划或进入回答。"),
    ("p", "Supervisor 的价值在于把“什么时候调用哪个 Agent”变成可控逻辑。首轮上传图像时，系统通常需要 VisionAgent；当用户询问原因、风险、维修建议或相似案例时，系统会考虑 KnowledgeAgent；当信息不足时，ClarificationAgent 可以生成澄清问题；当用户请求报告时，ReportAgent 整理完整诊断内容。这样的设计让任务执行更接近人工质检流程，而不是固定调用一个模型。"),
    ("table", "表4-1 Agent职责表", ["组件", "主要职责", "典型输出"], [
        ["Supervisor", "生成执行计划并决定调用哪些专家", "execution_plan、step_status"],
        ["VisionAgent", "调用视觉后端进行异常检测和定位", "anomalies、metadata、heatmap路径"],
        ["KnowledgeAgent", "调用RAG并生成缺陷/产品分析", "defect_analysis、object_analysis"],
        ["ClarificationAgent", "在信息不足时生成澄清问题", "pending_question"],
        ["ReportAgent", "汇总检测、知识和对话上下文", "诊断报告"],
        ["AnswerNode", "生成最终中文回答并处理失败状态", "answer、summary"],
    ]),
    ("h2", "视觉检测与PatchCore实现"),
    ("p", "视觉检测入口位于 app/tools/image_anomaly_detection.py。系统通过 resolve_visual_backend 解析用户选择或环境变量，决定使用 Qwen、专业 HTTP 服务或 PatchCore。Qwen 工具会构造包含输出格式、正常样本策略、误报控制和定位要求的 prompt，并要求模型返回 JSON。为了降低正常图误报，系统设置了 qwen_min_anomaly_score，并过滤分数过低或包含 reflection、shadow、edge、noise 等不确定描述的候选异常。"),
    ("p", "PatchCore 实现位于 app/tools/patchcore_detection.py。训练阶段读取某一类别 train/good 正常图像，提取 layer2 和 layer3 特征并合并为 patch embedding，随后保存 memory bank 和 metadata。推理阶段读取 memory bank，对输入图像计算 patch 到正常特征库的最近距离，将距离图上采样为原图大小的热力图，再根据阈值提取连通区域，输出 bbox、粗略位置、异常外观描述、heatmap、mask 和 overlay。"),
    ("h2", "RAG知识库与Few-shot实现"),
    ("p", "RAG 模块位于 app/rag。DatasetAnalyzer 负责按 MVTec 目录结构扫描样本，识别 train/good、test/good 和 test/<defect_type> 等路径；TextGenerator 用于生成样本文本描述；VectorStore 将文本或图像相关信息写入 Chroma；Retriever 负责查询相似案例并格式化为 prompt_context。后端提供同步建库、异步建库、查询和在线反馈入库接口，前端可以直接触发这些接口。"),
    ("p", "Few-shot 链路由 app/rag/fewshot.py 实现。FewShotSelector 会对 RAG 返回结果去重，并按照元数据中的 is_anomaly 或 anomaly_type 判断样本属于正常还是异常，再分别选择有限数量的 normal 和 anomaly 示例。FewShotPromptBuilder 将这些样本整理成紧凑文本块。该机制的定位是提示级参考，而不是参数训练；它的作用是让视觉或知识分析在判断时看到相似正常/异常案例。"),
    ("h2", "状态记忆与追问续接实现"),
    ("p", "状态模型定义在 app/memory/state.py。DetectionState 保存任务、上下文、日志、错误、检测结果、工具调用、Agent 输出、执行计划、执行事件、对话历史、对话摘要、用户回复、反思决策和报告请求等信息。为了降低状态理解成本，系统又将这些字段组织为 task_runtime、orchestration_runtime 和 domain_runtime 三个视图，分别对应任务输入、编排过程和领域结果。"),
    ("p", "追问续接依赖 checkpoint 和状态恢复。首轮检测完成后，用户如果继续提问，系统会根据 task_id 找回上一轮检测结果、共享上下文、RAG 信息和对话历史，再由 Supervisor 判断是否需要重新调用专家 Agent。长对话场景下，系统会将较早对话压缩为 conversation_summary，同时保留最近若干轮消息，避免上下文无限增长。"),
    ("h2", "前端交互与报告生成实现"),
    ("p", "前端页面位于 web/index.html，主要承担任务发起和结果观察。页面提供检测按钮、后端选择、PatchCore 类别提示、RAG 建库入口、RAG 查询入口、热力图预览、多 Agent 执行时间线和 Markdown 渲染能力。对于长任务排障，时间线能够展示 Agent 执行状态、步骤输出和错误信息，使用户知道系统当前卡在哪一步。"),
    ("p", "报告生成由 ReportAgent 和后端 /v1/generate_report 接口支持。与即时回答相比，报告更强调结构化和归档性，会综合检测结果、RAG 分析、MMAD 七任务结构、对话历史和上下文信息。AnswerNode 则负责生成交互式最终回答，并且在视觉检测失败时避免直接给出“正常”的错误结论。"),
    ("h1", "系统测试与实验分析"),
    ("h2", "测试环境与数据来源"),
    ("p", "系统测试应区分本地开发环境和 AutoDL 服务器环境。本地环境主要用于前端交互、接口调试和论文截图；AutoDL 环境主要用于 PatchCore 训练和数据集实验。项目依赖由 pyproject.toml 记录，核心包括 FastAPI、LangGraph、LangChain、Chroma、DashScope、Pillow、OpenCV、NumPy 和 PyTorch 相关库。数据集主要使用 MVTec AD，路径需以实际运行记录为准。"),
    ("table", "表5-1 测试环境记录表", ["环境项", "记录内容", "说明"], [
        ["本地开发环境", "Windows + Python 3.11或以上", "用于前端、接口和文档调试"],
        ["服务器环境", "Ubuntu/AutoDL + Python环境", "用于PatchCore训练和数据集实验"],
        ["数据集", "MVTec AD", "用于正常/异常样本和类别训练"],
        ["模型文件", "memory_bank.pt、metadata.json", "PatchCore每个类别对应一套模型文件"],
        ["向量库", "Chroma", "用于RAG样本检索"],
    ]),
    ("h2", "接口与功能测试"),
    ("p", "接口测试主要验证系统主链路是否可用。/health 用于确认后端启动；/v1/stream 用于验证流式检测和前端实时展示；/v1/chat 用于验证追问；/v1/generate_report 用于验证报告生成；/v1/rag/build/start 和 /v1/rag/build/status/{task_id} 用于验证 RAG 异步建库；/v1/rag/query 用于验证相似案例检索；/v1/patchcore/categories 用于验证 PatchCore 类别检查。"),
    ("p", "功能测试不应只看接口是否返回 200，还要看返回内容是否符合业务预期。例如正常图像不应被强行解释为异常，PatchCore 未训练类别应返回明确提示，追问时系统应使用上一轮上下文而不是重复旧答案，视觉检测失败时 AnswerNode 不应给出“正常”的误导性回答。"),
    ("h2", "PatchCore异常定位实验"),
    ("p", "PatchCore 实验建议从已训练类别开始，例如 bottle。训练时使用该类别 train/good 正常图像生成 memory bank，推理时分别选择 test/good 和 test/broken_small 等样本进行对比。记录内容包括类别、图像路径、阈值、是否输出异常区域、热力图路径、mask 路径、overlay 路径和人工观察结果。"),
    ("p", "根据前期运行记录可以看到，阈值选择会明显影响误报和漏报。阈值过低时，good 样本也可能出现异常区域；阈值提高后，good 样本误报会减少，但异常样本也可能需要重新观察是否仍能定位到缺陷。因此，本文测试章节应把 PatchCore 结果写成阈值调试和示例验证，不应直接写成完整性能评测结论。"),
    ("h2", "RAG检索与Few-shot实验"),
    ("p", "RAG 实验分为建库和查询两部分。建库时记录数据集路径、是否包含正常样本、处理数量、向量数量和任务状态。查询时输入产品类别、缺陷描述或用户问题，检查返回的相似案例是否与当前任务相关。Few-shot 实验则检查返回结果中是否同时包含正常参考和异常参考，以及这些参考是否被写入 metadata 或 prompt_context。"),
    ("p", "在论文中，RAG 实验更适合采用案例分析方式。例如，同一张图像在无 RAG 时只能得到较泛化描述；加入相似案例后，KnowledgeAgent 可以补充可能成因、风险提示和建议动作。这里需要强调的是，RAG 提供的是知识辅助和参考证据，不等同于自动证明模型判断一定正确。"),
    ("h2", "多Agent流程与追问测试"),
    ("p", "多 Agent 流程测试重点观察 Supervisor 是否根据任务意图选择合适节点。首轮图像检测应触发 VisionAgent；询问原因、风险、维修建议或相似案例时应触发 KnowledgeAgent；信息不足时应触发 ClarificationAgent；请求报告时应触发 ReportAgent。前端时间线和后端日志可用于记录每一步状态。"),
    ("p", "追问测试需要验证上下文处理。用户在首轮检测后继续提问“为什么会这样”“有没有使用 RAG”“重新分析一下”时，系统应基于同一个 task_id 读取上一轮状态，而不是把追问当成完全新的单轮任务。该测试可以直接反映状态记忆和 checkpoint 设计是否真正发挥作用。"),
    ("h2", "本章小结"),
    ("p", "本章从接口、前端、PatchCore、RAG、few-shot 和多 Agent 流程几个角度设计测试。由于部分量化指标需要更完整的数据集评测脚本和人工复核，本章当前以功能验证和案例实验为主。这样的写法更符合当前项目证据：系统已经具备可执行链路，但性能指标仍需要后续在更多类别和更多样本上系统统计。"),
    ("h1", "总结与展望"),
    ("h2", "工作总结"),
    ("p", "本文围绕工业异常检测中的定位、解释、追问和报告需求，设计并实现了一个基于多 Agent 与 RAG 的工业异常检测系统。系统以 FastAPI 作为服务入口，以 LangGraph 组织任务流程，以 Supervisor 协调不同专家 Agent，并将通用视觉模型、专业检测服务、PatchCore、本地 RAG、few-shot 上下文和 MMAD 结构化分析接入同一条诊断链路。"),
    ("p", "从工程实现看，系统已经形成了从前端上传图像、后端执行检测、RAG 检索相似案例、生成知识分析、展示热力图和时间线、继续追问、生成报告的闭环。与单次模型调用相比，该系统更强调过程可观察、上下文可续接和结果可解释，这也是本文工作的主要价值。"),
    ("h2", "系统不足"),
    ("p", "当前系统仍有不足。第一，PatchCore 依赖已训练类别，对未训练类别不能直接生成可靠热力图。第二，通用视觉模型仍可能受到成像环境和正常纹理干扰，需要更多正常样本和阈值策略辅助降低误报。第三，RAG 知识库的质量取决于数据集描述、案例数量和专家知识积累，目前仍以数据集样本和项目内描述为主。第四，测试章节尚未形成大规模量化评测，更多结论仍需要后续实验支撑。"),
    ("h2", "后续优化方向"),
    ("p", "后续可以从四个方面继续改进。首先，完善 PatchCore 多类别训练和阈值标定流程，为每个类别保存更可靠的推荐阈值。其次，引入更强的视觉基础模型或分割模型，提高细粒度异常定位能力。第三，扩充 RAG 知识库来源，将真实维修记录、工艺文档和专家标注案例纳入系统。第四，建立自动化评测脚本，对异常判别、缺陷定位、RAG 解释和 MMAD 七任务输出进行更系统的统计。"),
    ("h2", "本章小结"),
    ("p", "总体来看，本文并不是提出一个全新的异常检测算法，而是从系统工程角度把视觉检测、知识增强、多 Agent 编排和前端可视化组合起来。这个方向更贴近实际使用中的问题：检测结果不仅要能看，还要能解释、能追问、能复核、能形成报告。后续工作将继续围绕数据质量、模型能力和实验评估三个方面展开。"),
]


def paragraph_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def find_paragraph(doc: Document, predicate):
    for paragraph in doc.paragraphs:
        if predicate(paragraph):
            return paragraph
    raise RuntimeError("paragraph anchor not found")


def set_paragraph_text(paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_name(paragraph) -> str:
    try:
        return paragraph.style.name
    except Exception:
        return ""


def remove_between(start_para, stop_para, *, include_start: bool, include_stop: bool = False) -> None:
    body = start_para._element.getparent()
    children = list(body)
    start = children.index(start_para._element)
    stop = children.index(stop_para._element)
    delete_start = start if include_start else start + 1
    delete_stop = stop + 1 if include_stop else stop
    for element in children[delete_start:delete_stop]:
        body.remove(element)


def add_para_before(ref_para, text: str, style: str = BODY_STYLE, alignment=None, bold: bool = False):
    para = ref_para.insert_paragraph_before(text)
    try:
        para.style = style
    except Exception:
        pass
    if alignment is not None:
        para.alignment = alignment
    for run in para.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = bold
    return para


def add_heading_before(ref_para, text: str, level: int):
    style = H1_STYLE if level == 1 else H2_STYLE if level == 2 else H3_STYLE
    return add_para_before(ref_para, text, style=style)


def add_table_before(ref_para, caption: str, headers: list[str], rows: list[list[str]]) -> None:
    add_para_before(ref_para, caption, style=BODY_STYLE, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    doc = ref_para.part.document
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        header_cells[idx].text = header
        header_cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
    ref_para._p.addprevious(table._tbl)


def append_field_char(paragraph, field_type: str) -> None:
    run = OxmlElement("w:r")
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), field_type)
    run.append(fld_char)
    paragraph._p.append(run)


def append_instr_text(paragraph, instruction: str) -> None:
    run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run.append(instr)
    paragraph._p.append(run)


def replace_abstract(doc: Document) -> None:
    cn_title = find_paragraph(doc, lambda p: paragraph_text(p) == "摘　　要")
    cn_keywords = find_paragraph(doc, lambda p: paragraph_text(p).startswith("关键词："))
    remove_between(cn_title, cn_keywords, include_start=False, include_stop=False)
    for text in CN_ABSTRACT:
        add_para_before(cn_keywords, text, style=BODY_STYLE)

    en_title = find_paragraph(doc, lambda p: paragraph_text(p) == "Abstract")
    en_keywords = find_paragraph(doc, lambda p: paragraph_text(p).startswith("Keywords:"))
    remove_between(en_title, en_keywords, include_start=False, include_stop=False)
    for text in EN_ABSTRACT:
        add_para_before(en_keywords, text, style=BODY_STYLE)


def replace_toc(doc: Document) -> None:
    toc_title = find_paragraph(doc, lambda p: paragraph_text(p) == "目　　录")
    # The source draft cached a Word TOC field across many paragraphs. After
    # restructuring the body we rebuild the visible TOC entries and clear the
    # old dangling field markers on the title paragraph.
    set_paragraph_text(toc_title, "目　　录")
    append_field_char(toc_title, "begin")
    append_instr_text(
        toc_title,
        ' TOC \\o "2-3" \\h \\z \\t "标题 1,1,目录索引标题,1,目录索引加宽标题,1" ',
    )
    append_field_char(toc_title, "separate")
    first_body_heading = find_paragraph(
        doc,
        lambda p: style_name(p).lower() == H1_STYLE and paragraph_text(p) == "绪论",
    )
    remove_between(toc_title, first_body_heading, include_start=False, include_stop=False)
    last_toc_para = toc_title
    for level, text in TOC_ENTRIES:
        style = TOC1_STYLE if level == 1 else TOC2_STYLE if level == 2 else TOC3_STYLE
        last_toc_para = add_para_before(first_body_heading, text, style=style)
    append_field_char(last_toc_para, "end")


def replace_body(doc: Document) -> None:
    first_body_heading = find_paragraph(
        doc,
        lambda p: style_name(p).lower() == H1_STYLE and paragraph_text(p) == "绪论",
    )
    references = find_paragraph(
        doc,
        lambda p: paragraph_text(p) == "参考文献" and "toc" not in style_name(p).lower(),
    )
    remove_between(first_body_heading, references, include_start=True, include_stop=False)

    for item in BODY_BLOCKS:
        kind = item[0]
        if kind == "h1":
            add_heading_before(references, item[1], 1)
        elif kind == "h2":
            add_heading_before(references, item[1], 2)
        elif kind == "h3":
            add_heading_before(references, item[1], 3)
        elif kind == "p":
            add_para_before(references, item[1], style=BODY_STYLE)
        elif kind == "table":
            _, caption, headers, rows = item
            add_table_before(references, caption, headers, rows)
        else:
            raise ValueError(f"unknown block type: {kind}")


def force_black_text(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: rewrite_thesis_working_draft.py <docx_path>")
        return 2
    path = Path(sys.argv[1])
    if not path.exists():
        raise FileNotFoundError(path)

    doc = Document(str(path))
    replace_abstract(doc)
    replace_toc(doc)
    replace_body(doc)
    force_black_text(doc)
    doc.save(str(path))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
