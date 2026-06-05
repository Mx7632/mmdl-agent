from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


BODY_STYLE = "论文正文"
H1_STYLE = "heading 1"
H2_STYLE = "heading 2"


EXPANSIONS: dict[str, list[str]] = {
    "研究背景与意义": [
        "从工程应用角度看，工业异常检测的难点并不只在算法本身。生产线上的图像往往受到光照、相机角度、产品批次和背景夹具影响，同一类正常产品也可能存在细微差异。如果检测系统只给出一个分数，现场人员仍然需要再去判断这个分数意味着什么、是否需要停线、是否需要复检。因此，一个更贴近现场的系统，应当把模型结果翻译成可理解、可复核的诊断信息。",
        "本文项目正是在这个背景下展开。系统不是单纯追求替代人工，而是希望把人工经验、历史案例和视觉模型连接起来。对于明显异常，系统应当尽快给出位置和缺陷描述；对于不确定结果，系统应当保留追问入口；对于需要复盘的任务，系统应当能输出报告。这种设计虽然比单模型调用复杂一些，但更符合工业质检中“先发现、再解释、再复核、最后归档”的工作习惯。",
    ],
    "国内外研究现状": [
        "从已有研究看，工业异常检测大致经历了由传统特征方法到深度特征建模，再到多模态理解和知识增强的演进。MVTec AD 数据集推动了无监督工业异常检测评测的发展[1]，PatchCore、PaDiM、STFPM 等方法则说明，基于预训练特征的正常样本建模在工业检测中具有较强适用性[2-5]。这类方法的共同特点是重视局部特征差异，因而较适合缺陷定位。",
        "近两年，多模态大模型和视觉语言模型开始被用于工业异常理解。AnomalyGPT、MMAD、WinCLIP 等研究说明，缺陷检测正在从“模型输出一个异常分数”逐渐走向“模型理解异常并用语言解释异常”[7-8,14]。与此同时，RAG 和 Agent 技术的发展让系统可以在回答前主动检索外部知识，或根据任务状态调用工具[9-11]。这些研究给本文带来的启发是：异常检测系统应当同时关注视觉证据、知识证据和执行过程。",
    ],
    "现有问题分析": [
        "第一个问题是视觉检测结果与业务解释之间存在距离。很多异常检测算法能够输出热力图，却不会告诉用户这个热区可能对应划痕、污染还是破损，也不会说明它可能影响产品的哪一项功能。对于论文系统来说，如果只展示热力图，系统完整性不够；对于实际场景来说，如果没有文字解释，检测结果也很难被非算法人员快速理解。",
        "第二个问题是通用大模型的开放能力和工业检测的严谨性之间存在矛盾。通用模型擅长描述图像，但在正常样本和微弱异常之间容易“过度解释”。一张 good 样本如果存在反光或背景边缘，模型可能也会给出疑似异常。本文系统因此没有把 Qwen 等视觉模型的输出直接当成最终事实，而是在 prompt、阈值过滤、few-shot 参考和后续知识分析中增加约束。",
        "第三个问题是多轮任务缺少状态闭环。工业检测经常不是一次问答就结束，用户可能会继续问“为什么”“有没有相似案例”“重新分析一下”“生成报告”。如果系统无法记住上一轮检测结果，就会出现重复回答或脱离上下文的问题。本文把状态记忆、checkpoint 和对话摘要纳入设计，正是为了解决这类多轮任务的连续性问题。",
    ],
    "本文主要工作": [
        "在系统实现层面，本文工作并不是简单地把几个工具并列放在一起，而是围绕一个统一任务状态进行组织。DetectionTask 负责描述输入任务，DetectionState 负责保存运行过程，Supervisor 根据状态决定下一步调用哪个 Agent。这样设计后，视觉检测、RAG 检索、人工澄清和报告生成都可以共享同一个任务上下文。",
        "在论文贡献表述上，本文将重点放在工程集成和流程设计，而不是宣称提出新的基础模型。也就是说，本文的创新更接近“如何把已有视觉模型、PatchCore、RAG 和多 Agent 编排为一个可运行系统”。这种定位更符合当前项目证据，也能避免把尚未训练或尚未系统评测的能力写得过满。",
    ],
    "工业异常检测技术": [
        "工业异常检测通常可以分为图像级异常判别和像素级异常定位两类。图像级异常判别关注整张图是否存在缺陷，输出通常是正常或异常；像素级定位关注异常区域在哪里，输出可能是热力图、掩码或边界框。本文系统同时需要这两类结果，因为前端展示和报告生成都需要知道异常位置，而 KnowledgeAgent 的缺陷分析也需要依赖位置、形态和严重程度等信息。",
        "从数据特点看，工业异常检测很难像普通分类任务那样收集均衡数据。一方面，真实异常样本出现频率低；另一方面，异常类型可能不断变化。因此，许多方法选择只利用正常样本建模，再把偏离正常模式的样本视为异常。PatchCore、PaDiM 和 FastFlow 等方法都可以看作对正常特征分布的不同建模方式[2-3,18]。这也解释了为什么本文在本地工具中优先接入 PatchCore。",
    ],
    "PatchCore异常定位方法": [
        "PatchCore 对本文项目特别有价值，是因为它的输入输出形式与系统需求比较吻合。训练阶段只需要同类别的正常图片，推理阶段能够返回异常分数和空间热力图。系统可以把热力图进一步转成 mask、overlay、bbox 和中文位置描述，再交给前端展示或报告模块引用。与只返回一句自然语言的视觉模型相比，PatchCore 的结果更适合做可视化证据。",
        "但 PatchCore 的边界也必须写清楚。它依赖类别内正常样本建立 memory bank，因此不能自然泛化到完全没有训练过的新类别。如果把未训练类别也强行送入 PatchCore，结果很可能没有意义。当前项目通过 /v1/patchcore/categories 接口和后端类别检查逻辑解决这个问题：只有当类别存在 memory_bank.pt 和 metadata.json 时才允许使用 PatchCore，否则提示用户先训练该类别或切换到 Qwen 后端。",
        "此外，PatchCore 阈值并不是一个放之四海皆准的常数。前期实验中，较低阈值会让 good 样本也出现热区，较高阈值则可能减少误报但增加漏报风险。因此，第五章实验不应把单一阈值结果写成确定结论，而应将其视为类别相关的调参过程。这种保守写法更接近当前项目的真实状态。",
    ],
    "多模态大模型与MMAD任务": [
        "多模态大模型在本文中承担的是开放视觉理解和语言表达角色。它能够根据图像和提示词给出异常描述，也能够在用户追问时解释可能原因。不过，通用模型缺少工业现场的专门约束，因此系统需要把输出限定为 JSON，并在 prompt 中明确要求不要把正常纹理、反光、阴影和边缘强行识别为异常。这些约束对应项目中的 Qwen 工具实现。",
        "MMAD 的意义在于提供了一种更完整的任务拆分视角[8]。如果只做 anomaly discrimination，系统只需要判断是否异常；如果加入 defect localization，系统还要说明位置；如果加入 defect analysis 和 object analysis，系统就必须进一步解释缺陷可能影响什么。本文将这些任务作为结构化诊断的组织方式，而不是把它写成已经完全复现 MMAD 论文的数据生成流水线。",
    ],
    "RAG与少样本上下文": [
        "RAG 在本文中有两类使用时机。第一类发生在视觉检测前或视觉检测过程中，系统尝试从知识库中检索相似正常样本和异常样本，再把这些案例作为 few-shot 参考注入视觉 prompt。这样做的目的不是让检索结果替代模型判断，而是让模型看到“同类正常外观”和“相似缺陷外观”的对照。",
        "第二类发生在 KnowledgeAgent 执行时。当用户关心缺陷原因、风险提示、维修建议或相似案例时，系统通过 RAG 查询知识库，并将结果整理成 defect_analysis 和 object_analysis。这里的知识增强更像是把历史样本和规则经验放到模型面前，让回答不再完全依赖模型自身记忆。与直接问大模型相比，这种方式更容易解释回答依据。",
        "需要特别说明的是，本文中的 few-shot 机制属于提示级少样本，不涉及参数更新。FewShotSelector 主要根据检索结果中的 metadata 判断样本是 normal 还是 anomaly，并控制正常样本与异常样本的数量。它的工程价值在于把 RAG 结果变成更可读、更可注入的上下文，而不是训练一个新的少样本分类器。",
    ],
    "LangGraph与多Agent协作": [
        "多 Agent 并不是为了让系统显得复杂，而是为了把不同问题交给不同角色处理。视觉检测、知识解释、人工澄清和报告生成使用的是不同上下文、不同输入输出和不同失败处理方式。如果全部写在一个节点里，后续调试会非常困难。拆分成 Agent 后，系统可以在时间线中展示每个 Agent 的执行状态，也可以单独测试某个 Agent 的输入输出是否正确。",
        "LangGraph 的优势在于它把流程节点和状态结合起来。本文系统中，节点之间不是简单顺序执行，而是会根据执行计划、用户是否需要澄清、反思结果和报告请求进行条件跳转。例如 self_reflect 可以决定是否重试，wait_user 可以等待人工输入，answer 可以根据 report_requested 决定是否进入 report 节点。这种图式编排比固定流水线更适合多轮诊断任务。",
    ],
    "应用场景与使用角色": [
        "在实际使用中，系统并不要求所有用户都理解底层模型。检测人员更关心结果是否清楚，最好能直接看到异常区域；工程技术人员更关心异常原因和处理建议；系统维护人员则需要知道 RAG 是否建库成功、PatchCore 类别是否训练、后端接口是否正常。论文将这些人群写成“使用场景角色”，而不是账号权限角色，能够避免虚构登录、权限和审计模块。",
        "这类角色划分也有助于解释为什么系统需要前端可视化。对于检测人员来说，热力图和 overlay 比后端 JSON 更直观；对于维护人员来说，Agent 时间线和错误提示能帮助定位任务卡在哪一步；对于论文实验来说，RAG 构建状态、PatchCore 类别状态和流式事件都是可观察证据。前端不是附属页面，而是系统可解释性的一部分。",
    ],
    "功能需求分析": [
        "图像检测需求要求系统至少完成三件事：接收用户上传图片、选择合适视觉后端、返回结构化检测结果。结构化结果中应包含 status、anomalies、summary 和 metadata 等字段。这样做的好处是后续模块不需要解析大段自由文本，而可以直接读取异常列表、位置、热力图路径和后端信息。",
        "知识增强需求要求系统能够在合适时机调用 RAG，而不是每次都强制检索。比如用户只问“这张图有没有异常”时，系统可以先完成视觉检测；当用户进一步问“原因是什么”“有没有相似案例”时，KnowledgeAgent 再进入流程。这样的设计兼顾响应速度和知识支撑，也符合当前 Supervisor 规划节点的职责。",
        "追问与报告需求体现了系统的多轮特征。用户在首轮检测后继续提问时，系统必须复用上一轮 task_id 对应的状态；当用户请求报告时，系统应当把视觉结果、知识分析、MMAD 结构和对话历史合并起来。若没有状态记忆，报告生成就只能重新组织当前文本，无法真正反映前面发生过什么。",
    ],
    "非功能需求分析": [
        "可维护性体现在代码分层上。接口层只负责接收请求和返回响应，图编排层负责控制流程，多 Agent 层负责专家职责，工具层封装外部模型或本地算法，RAG 层负责知识库构建和检索，状态层负责上下文保存。这种分层虽然增加了一些文件数量，但能降低模块之间互相牵扯的风险。",
        "可观测性是本文系统的重要非功能需求。工业检测任务一旦失败，用户需要知道是图像上传失败、视觉模型返回格式错误、RAG 检索为空、PatchCore 类别未训练，还是 Supervisor 路由出现循环。系统通过日志 trace_id、SSE 流式事件、前端时间线和执行事件列表提供观察入口。论文中可以把这一点写成系统设计目标，而不应写成成熟监控平台。",
        "可扩展性主要体现在后端选择和模块替换上。视觉后端可以从 Qwen 切换到 PatchCore 或专业 HTTP 服务；RAG 知识库可以继续扩充样本和描述；Agent 节点也可以根据论文后续工作继续增加。当前项目还没有证明大规模并发能力，因此可扩展性应理解为软件结构可扩展，而不是部署规模已经可扩展。",
    ],
    "系统总体架构设计": [
        "六层架构之间并不是简单上下堆叠，而是围绕任务状态协作。用户交互层产生请求，API 层将请求转换为 DetectionTask，图编排层将任务推进到不同节点，多 Agent 层决定具体执行者，工具与知识层提供检测和检索能力，状态记忆层则贯穿整个流程。状态记忆层不是最后才使用，而是在每次规划、执行、合并和回答时都可能被读取或更新。",
        "这种架构还有一个好处：它把“模型能力”和“系统能力”区分开。模型能力指 Qwen、PatchCore 或专业检测服务本身能识别什么；系统能力指项目如何组织输入、如何选择工具、如何处理失败、如何保存上下文、如何展示结果。本文论文的重点应放在系统能力上，因为这部分最能体现本科毕业设计的工程实现工作。",
        "在图 3-1 的架构图中，建议把用户交互层、API 与任务运行层、图编排层、多 Agent 协作层、工具与知识层、状态记忆层画成纵向或分组结构。不要照搬通用 PaaS 图中的多租户、安全中心、容器编排等内容，因为这些并不是当前项目已经实现的重点。论文图应当服务真实项目，而不是服务概念完整性。",
    ],
    "主运行流程设计": [
        "主流程的关键在于 Supervisor 不是一次性决定全部结果，而是根据状态逐步推进。首轮可能只需要视觉检测；如果视觉结果不确定，流程可能进入澄清；如果用户询问知识性问题，则再调用 RAG；如果报告请求被置为 true，则进入 report 节点。这种运行方式使系统能够处理不同复杂度的任务，而不是把所有任务都拉成同一条固定流水线。",
        "此外，流程设计还必须考虑失败路径。比如视觉后端没有配置 API key，PatchCore 类别没有训练，RAG 知识库为空，模型返回不是 JSON，或者图循环没有到达停止条件。这些情况如果没有统一处理，前端就只能看到模糊的错误。本文系统通过异常类型、状态字段、反思节点和错误事件记录，把失败变成可展示、可诊断的信息。",
    ],
    "API接口与任务模型实现": [
        "接口层的设计重点是把外部请求转换成统一内部任务。/v1/stream 和 /v1/detect 都会读取 multipart/form-data，如果存在 image 文件，就把图像转为 base64，并写入 parameters。这样，后续工具无需关心文件上传细节，只需要从 DetectionTask 中读取 image_base64、image_mime 和 detector_params。",
        "RAG 接口单独设计，是因为知识库构建和检测任务不是同一种生命周期。建库可能是一个较长任务，需要 start/status 两个接口配合；查询则是短任务，可以直接返回结果和 prompt_context。把 RAG 接口前置到前端之后，用户不必每次通过命令行建库，这也让论文系统更像一个完整应用，而不是只面向开发者的脚本集合。",
    ],
    "图编排与Supervisor实现": [
        "Supervisor 计划阶段的输出可以看作任务清单，而执行阶段则把这些清单交给具体专家节点。执行完成后，merge 节点统一把结果写入共享上下文。这样做可以避免不同 Agent 各自修改最终答案，造成结果混乱。最终回答仍由 AnswerNode 汇总生成，从而保持用户看到的回答风格和错误处理规则一致。",
        "self_reflect 节点用于检查流程是否需要重试或澄清。前期开发中曾出现视觉分析失败但最终回答仍显示正常的问题，这说明反思节点和回答节点必须正确传递失败状态。修复后的逻辑强调：如果关键视觉步骤失败，系统不应把失败当作无异常，而应明确告诉用户检测失败并给出可操作提示。这一设计可以在论文中作为可靠性处理案例。",
    ],
    "视觉检测与PatchCore实现": [
        "Qwen 后端的实现重点是结构化输出和误报过滤。系统在 prompt 中要求模型返回 JSON，并明确正常样本应返回空 anomalies。随后，filter_qwen_anomalies 会根据分数和描述文本过滤弱异常、反光、阴影、边缘、纹理和噪声等不确定候选。这样做不能保证完全没有误报，但能把通用视觉模型的开放描述收束到更适合工业检测的输出格式。",
        "PatchCore 后端的实现更偏算法工具。训练时，脚本会遍历 train/good 图像并提取特征；推理时，系统把 patch 距离转换为热力图，再用连通域方法提取异常区域。输出中的 location、appearance、severity_hint 和 description 并不是模型直接生成，而是根据 bbox、面积比例、热区强度和形状规则推导出来。这种处理让热力图结果更容易进入报告文字。",
        "两类视觉后端的互补关系应写清楚。Qwen 更适合开放类别和自然语言描述，但可能误报；PatchCore 更适合已训练类别和可视化定位，但需要正常样本训练。前端允许用户选择后端，后端也会在 metadata 中记录 selected_backend，这使实验时能够区分不同后端的行为。",
    ],
    "RAG知识库与Few-shot实现": [
        "RAG 建库并不是简单保存文件路径。DatasetAnalyzer 需要识别样本类别、是否异常、缺陷类型和可能的 ground truth 路径；TextGenerator 将这些结构信息转为可检索描述；VectorStore 再把描述和元数据写入 Chroma。这样，后续查询时不仅能返回文本，还能返回 category、anomaly_type、is_anomaly、image_path 等用于 few-shot 选择的字段。",
        "FewShotSelector 的实现刻意保持简单：先去重，再按 metadata 判断正常或异常，最后限制 normal 和 anomaly 数量。这样的设计虽然不如复杂学习模型强，但符合当前项目阶段。它不需要额外训练，也不需要用户提供大量标注；只要 RAG 返回结果中包含可靠元数据，就能构造正常/异常对照样本。",
        "KnowledgeAgent 使用 RAG 的方式更偏解释。它会根据用户问题或异常列表构造查询文本，检索相似案例，再生成 possible_causes、risk_notes、repair_actions 和 similar_cases。对于产品分析，系统还会结合 object_context 和对象知识生成 object_profile、component_scope 和 functional_impact。这些字段让报告不只是“发现异常”，还能够进一步说明异常可能影响什么。",
    ],
    "状态记忆与追问续接实现": [
        "DetectionState 中的字段较多，但可以理解为三类信息。第一类是任务信息，例如用户问题、上传图像、报告请求和对话历史。第二类是编排信息，例如执行计划、当前 Agent、step 状态、重试次数和反思决策。第三类是领域信息，例如检测结果、工具输出、共享上下文、RAG 结果和 Agent trace。三类信息分开后，开发者更容易定位问题。",
        "对话压缩的必要性来自上下文窗口限制。如果用户连续追问十几轮，把所有消息完整塞入模型不仅成本高，也可能让模型被早期无关内容干扰。系统通过 conversation_summary 保存较早对话的摘要，同时保留最近对话，实现“记住大意但不无限增长”。这种设计与工业排障很像：保留关键结论和最近操作，而不是把所有细节都重复一遍。",
        "追问链路中最重要的是 task_id。只要 task_id 保持一致，系统就能从 checkpoint 中恢复上一轮状态；如果 task_id 变化，系统就可能把追问当成新任务。论文中可以把这一点作为前后端协作要求写清楚，因为很多所谓“追问失败”并不是模型不会回答，而是状态没有被正确绑定。",
    ],
    "前端交互与报告生成实现": [
        "前端页面承担了几个调试价值很高的功能。第一，它能展示流式输出，用户不需要等到所有节点执行完才看到状态。第二，它能展示多 Agent 时间线，便于判断是 VisionAgent、KnowledgeAgent 还是 ReportAgent 出现问题。第三，它能预览 PatchCore 生成的 heatmap、mask 和 overlay，让异常定位不再只是后端路径。第四，它能触发 RAG 建库和查询，降低使用门槛。",
        "报告生成的目标不是替代论文实验，而是服务工业诊断闭环。即时回答适合交互，报告适合复核和归档。ReportAgent 会复用任务状态中的视觉结果、RAG 分析、MMAD 结构和对话历史，因此报告内容比单轮回答更完整。后续如果要加入 PDF 或 Word 报告导出，也可以在这一模块基础上扩展。",
    ],
    "测试环境与数据来源": [
        "测试环境记录应当尽量写成可复现形式。对于本地环境，需要记录操作系统、Python 版本、依赖安装命令、后端启动命令和前端访问地址；对于 AutoDL 环境，需要记录 Ubuntu 版本、GPU 型号、PyTorch/CUDA 兼容性、数据集路径和模型保存路径。前期训练中遇到 RTX 5090 与 PyTorch CUDA 架构不兼容的问题，这也说明环境记录不是形式化内容，而是实验复现的重要依据。",
        "数据来源方面，MVTec AD 适合作为 PatchCore 训练和推理验证数据集，因为其目录结构清晰，包含 train/good、test/good、test/<defect_type> 和 ground_truth。MMAD 更适合用于多模态任务思想和结构化输出设计，不应简单写成已经完整训练了 MMAD 模型。论文中要区分数据集用于训练、用于知识库、用于结构验证这几种不同用途。",
    ],
    "接口与功能测试": [
        "功能测试可以采用“输入、操作、预期结果、实际结果、是否通过”的表格方式记录。例如，上传一张已训练类别的 bottle 图像并选择 PatchCore，预期结果是返回 heatmap_path 和 overlay_path；上传未训练类别时，预期结果是返回类别未训练提示；构建 RAG 时，预期结果是状态从 running 变为 success；追问上一轮任务时，预期结果是回答引用已有检测结果。",
        "接口测试还要关注错误响应。一个系统是否好用，不只看成功路径，也看失败路径是否清楚。比如缺少 task_id、上传图像为空、parameters 不是合法 JSON、RAG task_id 不存在、PatchCore memory bank 缺失，这些都应返回明确错误。论文可以选择几类典型错误写入测试用例，说明系统具备基本异常处理能力。",
    ],
    "PatchCore异常定位实验": [
        "PatchCore 实验建议分两阶段写。第一阶段是训练验证，记录某一类别训练命令、数据集路径、image_size、backbone、是否使用预训练权重、memory_bank_size 和 recommended_threshold。第二阶段是推理验证，对 good 和 defective 样本分别运行 patchcore_eval.py，记录 threshold、anomaly_score、异常区域数量和可视化输出路径。",
        "对于误报问题，论文中不应回避。事实上，good 样本在较低阈值下被判为异常，是异常检测系统中很常见的阈值敏感问题。把这一点写出来反而更真实：系统已经能训练和推理，但阈值仍需结合类别、正常样本分布和业务容忍度调整。后续可以通过更多 good 样本统计分布，确定更稳健的推荐阈值。",
    ],
    "RAG检索与Few-shot实验": [
        "RAG 实验可以设计两组对比。第一组只验证检索质量：给定 category 和缺陷描述，观察返回样本是否来自相同或相近类别，metadata 是否包含 anomaly_type 和 is_anomaly。第二组验证 few-shot 注入：观察 VisionAgent 或 KnowledgeAgent 的 metadata 中是否记录 few_shot_examples 和 few_shot_context，以及正常/异常参考样本是否同时存在。",
        "解释效果评估可以采用人工评分，但评分表应保守设计。例如从相似案例相关性、原因解释合理性、建议动作具体性、是否出现幻觉四个维度打分。由于当前项目没有大量专家标注，评分只能作为小规模案例分析，不能写成大规模统计结论。这样的实验安排既能体现 RAG 价值，也不会超出证据边界。",
    ],
    "多Agent流程与追问测试": [
        "多 Agent 测试的核心是路由正确性。可以构造几类用户问题：上传图像并问是否异常，预期触发 VisionAgent；问缺陷原因，预期触发 KnowledgeAgent；问“生成报告”，预期触发 ReportAgent；当结果不确定时，预期进入 ClarificationAgent 或 wait_user。前端时间线和后端 execution_events 可以作为判断依据。",
        "追问测试需要特别关注“重复旧答案”问题。如果用户追问后系统没有重新规划，而是直接返回上一轮 summary，就说明 chat 链路没有正确进入 Supervisor 或没有把问题写入状态。论文可以把该问题作为系统迭代过程中的一个典型修复点：通过状态恢复、is_continue 标记和聊天接口逻辑，保证追问能够参与新一轮推理。",
    ],
    "工作总结": [
        "从论文完成度看，本文系统已经覆盖了一个工业异常检测智能诊断系统的主要链路。它不是只展示一个算法脚本，也不是只搭一个静态前端，而是把后端接口、图编排、多 Agent、视觉检测、RAG、PatchCore、状态记忆和前端可视化连在一起。这种完整链路是本科毕业设计中比较重要的工程证明。",
        "同时，本文保持了比较清晰的边界：系统实现了可执行流程，但并没有声称超过某个先进算法；系统接入了 RAG 和 few-shot，但没有声称完成模型微调；系统支持 PatchCore 热力图，但没有声称适用于所有未训练类别。这样的总结方式更稳，也更经得起代码和实验材料核查。",
    ],
    "系统不足": [
        "系统不足主要来自三个方面。第一，模型侧仍依赖外部能力，Qwen 的视觉判断和 PatchCore 的类别训练都会影响最终结果。第二，知识侧仍依赖数据集描述和案例质量，如果 RAG 中相似案例不足，KnowledgeAgent 的解释也会变得笼统。第三，实验侧仍缺少大规模自动化指标统计，目前更适合写成功能验证和案例分析。",
        "此外，当前前端虽然已经能展示多 Agent 时间线和热力图，但还可以继续增强。例如，可以把 RAG 命中的正常/异常参考样本直接展示在检测结果旁边，让用户看到模型参考了哪些样本；也可以在报告中自动插入 heatmap 和 overlay，使报告更像真实质检记录。这些都属于后续工程优化方向。",
    ],
    "后续优化方向": [
        "后续第一项工作是完善评测体系。可以为 MVTec 每个类别编写批量评估脚本，统计图像级 AUROC、像素级 AUROC、PRO 或人工复核结果，并把不同阈值下的误报率和漏报率画成表格。这样，第五章就能从当前的功能验证升级为更完整的实验分析。",
        "第二项工作是增强知识库质量。当前 RAG 更多来自数据集样本和自动描述，后续可以加入真实维修记录、工艺说明、专家标注和缺陷成因知识。知识库越接近真实工业场景，KnowledgeAgent 的回答就越不容易空泛。第三项工作是探索轻量微调，例如 LoRA、Adapter 或类别原型更新，使系统不仅能提示级适配，还能逐步学习新类别。",
    ],
}


REFERENCES = [
    "[1] Bergmann P, Fauser M, Sattlegger D, et al. MVTec AD -- A Comprehensive Real-World Dataset for Unsupervised Anomaly Detection[C]//Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition. 2019: 9592-9600.",
    "[2] Roth K, Pemula L, Zepeda J, et al. Towards Total Recall in Industrial Anomaly Detection[C]//Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition. 2022: 14318-14328.",
    "[3] Defard T, Setkov A, Loesch A, et al. PaDiM: a Patch Distribution Modeling Framework for Anomaly Detection and Localization[C]//Pattern Recognition. ICPR International Workshops and Challenges. Cham: Springer, 2021: 475-489.",
    "[4] Zavrtanik V, Kristan M, Skočaj D. DRAEM -- A Discriminatively Trained Reconstruction Embedding for Surface Anomaly Detection[C]//Proceedings of the IEEE/CVF International Conference on Computer Vision. 2021: 8330-8339.",
    "[5] Wang G, Han S, Ding E, et al. Student-Teacher Feature Pyramid Matching for Anomaly Detection[C]//Proceedings of the British Machine Vision Conference. 2021.",
    "[6] Batzner K, Heckler L, König R. EfficientAD: Accurate Visual Anomaly Detection at Millisecond-Level Latencies[C]//Proceedings of the IEEE/CVF Winter Conference on Applications of Computer Vision. 2024: 127-137.",
    "[7] Gu Z, Zhu B, Zhu G, et al. AnomalyGPT: Detecting Industrial Anomalies Using Large Vision-Language Models[C]//Proceedings of the AAAI Conference on Artificial Intelligence. 2024.",
    "[8] Jiang X, Li J, Deng H, et al. MMAD: The First-Ever Comprehensive Benchmark for Multimodal Large Language Models in Industrial Anomaly Detection[EB/OL]. arXiv:2410.09453, 2024.",
    "[9] Lewis P, Perez E, Piktus A, et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks[C]//Advances in Neural Information Processing Systems. 2020, 33: 9459-9474.",
    "[10] Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models[C]//International Conference on Learning Representations. 2023.",
    "[11] Wu Q, Bansal G, Zhang J, et al. AutoGen: Enabling Next-Gen LLM Applications via Multi-Agent Conversation Framework[EB/OL]. arXiv:2308.08155, 2023.",
    "[12] Li C L, Sohn K, Yoon J, et al. CutPaste: Self-Supervised Learning for Anomaly Detection and Localization[C]//Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition. 2021: 9664-9674.",
    "[13] Radford A, Kim J W, Hallacy C, et al. Learning Transferable Visual Models from Natural Language Supervision[C]//International Conference on Machine Learning. 2021: 8748-8763.",
    "[14] Jeong J, Zou Y, Kim T, et al. WinCLIP: Zero-/Few-Shot Anomaly Classification and Segmentation[C]//Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition. 2023: 19606-19616.",
    "[15] Liu H, Li C, Wu Q, et al. Visual Instruction Tuning[C]//Advances in Neural Information Processing Systems. 2023, 36: 34892-34916.",
    "[16] Kirillov A, Mintun E, Ravi N, et al. Segment Anything[C]//Proceedings of the IEEE/CVF International Conference on Computer Vision. 2023: 4015-4026.",
    "[17] Deng H, Li X. Anomaly Detection via Reverse Distillation from One-Class Embedding[C]//Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition. 2022: 9737-9746.",
    "[18] Yu J, Zheng Y, Wang X, et al. FastFlow: Unsupervised Anomaly Detection and Localization via 2D Normalizing Flows[EB/OL]. arXiv:2111.07677, 2021.",
    "[19] Liu Z, Zhou Y, Xu Y, et al. SimpleNet: A Simple Network for Image Anomaly Detection and Localization[C]//Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition. 2023: 20402-20411.",
    "[20] Zavrtanik V, Kristan M, Skočaj D. Reconstruction by Inpainting for Visual Anomaly Detection[J]. Pattern Recognition, 2021, 112: 107706.",
    "[21] Schlegl T, Seeböck P, Waldstein S M, et al. Unsupervised Anomaly Detection with Generative Adversarial Networks to Guide Marker Discovery[C]//International Conference on Information Processing in Medical Imaging. 2017: 146-157.",
    "[22] Kingma D P, Welling M. Auto-Encoding Variational Bayes[EB/OL]. arXiv:1312.6114, 2013.",
    "[23] He K, Zhang X, Ren S, et al. Deep Residual Learning for Image Recognition[C]//Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition. 2016: 770-778.",
    "[24] Dosovitskiy A, Beyer L, Kolesnikov A, et al. An Image is Worth 16x16 Words: Transformers for Image Recognition at Scale[C]//International Conference on Learning Representations. 2021.",
    "[25] Johnson J, Douze M, Jégou H. Billion-Scale Similarity Search with GPUs[J]. IEEE Transactions on Big Data, 2021, 7(3): 535-547.",
    "[26] Malkov Y A, Yashunin D A. Efficient and Robust Approximate Nearest Neighbor Search Using Hierarchical Navigable Small World Graphs[J]. IEEE Transactions on Pattern Analysis and Machine Intelligence, 2020, 42(4): 824-836.",
    "[27] Johnson J M, Khoshgoftaar T M. Survey on Deep Learning with Class Imbalance[J]. Journal of Big Data, 2019, 6(1): 27.",
    "[28] Ramírez S. FastAPI Documentation[EB/OL]. https://fastapi.tiangolo.com/.",
    "[29] LangChain. LangGraph: Multi-Agent Workflows[EB/OL]. https://www.langchain.com/blog/langgraph-multi-agent-workflows, 2024.",
    "[30] Chroma. Chroma Documentation: Collections and Embeddings[EB/OL]. https://docs.trychroma.com/.",
]


def paragraph_text(paragraph) -> str:
    return paragraph.text.strip()


def style_name(paragraph) -> str:
    try:
        return paragraph.style.name.lower()
    except Exception:
        return ""


def add_para_before(ref_para, text: str, style: str = BODY_STYLE):
    para = ref_para.insert_paragraph_before(text)
    try:
        para.style = style
    except Exception:
        pass
    for run in para.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return para


def find_body_heading(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph_text(paragraph) == text and style_name(paragraph) in {H1_STYLE, H2_STYLE}:
            return paragraph
    raise RuntimeError(f"heading not found: {text}")


def find_next_body_heading(doc: Document, heading_para):
    paragraphs = doc.paragraphs
    start = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph._element is heading_para._element
    )
    for paragraph in paragraphs[start + 1 :]:
        if style_name(paragraph) in {H1_STYLE, H2_STYLE} or paragraph_text(paragraph) == "参考文献":
            return paragraph
    raise RuntimeError(f"next heading not found after: {paragraph_text(heading_para)}")


def expand_sections(doc: Document) -> None:
    for heading, paragraphs in EXPANSIONS.items():
        heading_para = find_body_heading(doc, heading)
        stop_para = find_next_body_heading(doc, heading_para)
        for text in paragraphs:
            add_para_before(stop_para, text)


def replace_references(doc: Document) -> None:
    refs = None
    ack = None
    for paragraph in doc.paragraphs:
        text = paragraph_text(paragraph)
        if text == "参考文献" and "toc" not in style_name(paragraph):
            refs = paragraph
        if text == "致　　谢" and refs is not None and "toc" not in style_name(paragraph):
            ack = paragraph
            break
    if refs is None or ack is None:
        raise RuntimeError("reference or acknowledgement anchor not found")

    body = refs._element.getparent()
    children = list(body)
    start = children.index(refs._element)
    stop = children.index(ack._element)
    for element in children[start + 1 : stop]:
        body.remove(element)
    for item in REFERENCES:
        add_para_before(ack, item)


def force_black(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: expand_thesis_working_draft.py <docx_path>")
        return 2
    path = Path(sys.argv[1])
    doc = Document(str(path))
    expand_sections(doc)
    replace_references(doc)
    force_black(doc)
    doc.save(str(path))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
