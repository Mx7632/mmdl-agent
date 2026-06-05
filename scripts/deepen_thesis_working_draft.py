from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


BODY_STYLE = "论文正文"
H1_STYLE = "heading 1"
H2_STYLE = "heading 2"


DEEPEN_BLOCKS: dict[str, list[str]] = {
    "研究背景与意义": [
        "进一步说，工业异常检测系统的价值并不只取决于模型是否先进，还取决于它能否进入实际工作流。一个算法在论文数据集上表现很好，但如果没有可上传图片的入口、没有可解释的结果、没有错误处理、没有追问能力，现场人员仍然很难使用。本文把前端、接口、任务编排、知识库和报告放在同一系统中，就是希望把算法能力转化为一个能被人操作、能被人理解、也能被人复盘的应用形态。",
    ],
    "国内外研究现状": [
        "从研究趋势看，异常检测正在从单一模型评测转向更完整的任务理解。一方面，PatchCore、FastFlow、Reverse Distillation、SimpleNet 等方法持续优化检测和定位效果[17-19]；另一方面，LLaVA、SAM、CLIP 等基础模型推动了跨任务视觉语义能力的发展[13,15-16]。本文系统并没有重新训练这些基础模型，而是把它们代表的能力边界纳入系统设计：本地算法负责稳定热力图，通用模型负责开放描述，RAG 负责知识补充，多 Agent 负责流程组织。",
    ],
    "应用场景与使用角色": [
        "以一次典型检测为例，检测人员首先上传产品图片并选择检测后端。系统返回异常位置后，检测人员可能只需要判断是否复检；工程技术人员则会继续追问缺陷成因和影响；维护人员会关注当前类别是否已经训练、知识库是否构建、接口是否返回异常。这些行为在前端看起来只是几个按钮和输入框，但在后端会对应不同的接口、状态字段和 Agent 路由。需求分析必须把这些行为拆开，否则后续总体设计就会显得没有依据。",
        "本系统目前没有实现登录、角色权限和审批流程，因此论文不应把这些使用角色写成权限模型。更合理的写法是把它们作为系统服务对象，用来说明为什么需要结果展示、知识解释、状态追踪和报告生成。这样既能体现需求分析的完整性，也不会引入代码中不存在的功能。",
    ],
    "功能需求分析": [
        "从输入输出角度看，图像检测功能的输入包括 task_id、asset_id、question、image、parameters 等信息，输出包括状态、异常列表、摘要、回答和元数据。元数据在系统中很重要，它记录了 selected_backend、localization_available、heatmap_path、few_shot_context 等信息。没有元数据，前端就很难判断该展示热力图、展示 RAG 参考，还是提示用户检查配置。",
        "RAG 建库功能与检测功能不同，它更像一个后台准备过程。用户需要指定数据集根目录，系统扫描样本、生成描述、写入向量库，并通过 build/status 接口返回进度。这个功能的存在说明本文系统不只依赖在线大模型，而是主动构建了项目自己的知识来源。对毕业论文而言，RAG 建库入口也能证明知识库不是停留在概念层，而是有可调用的工程接口。",
        "报告生成功能可以看作检测链路的最终沉淀。检测结果适合即时查看，但报告需要把异常位置、检测证据、相似案例、原因分析、风险提示和建议动作组织起来。这个需求对应 ReportAgent 和 /v1/generate_report 接口，也对应论文中“系统不仅能回答，还能形成诊断记录”的设计目标。",
    ],
    "非功能需求分析": [
        "可靠性需求主要体现在失败不误导。工业检测系统宁可告诉用户“当前检测失败，需要检查模型配置或类别训练状态”，也不应该把失败包装成“未发现异常”。因此，系统在视觉工具、Supervisor 合并、self_reflect 和 AnswerNode 中都需要保留失败状态。论文中可以把这一点作为非功能需求说明：系统应当保证错误可见，而不是保证永远成功。",
        "易用性需求体现在前端减少命令行依赖。早期 RAG 建库、PatchCore 训练和推理都可以通过脚本完成，但如果论文演示完全依赖命令行，系统应用性会显得不足。当前前端增加了 RAG 建库和查询入口、PatchCore 类别状态、热力图预览和时间线展示，正是为了让用户在浏览器中完成更多操作。",
    ],
    "系统总体架构设计": [
        "用户交互层与 API 层之间采用 HTTP 接口连接。对于普通检测，前端以 multipart/form-data 上传图片；对于 RAG 查询，前端发送 JSON 请求；对于流式检测，后端返回 text/event-stream。不同接口形式对应不同交互需求：图像上传需要表单，知识查询适合 JSON，长任务进度适合 SSE。这种接口设计不是随意选择，而是和用户体验直接相关。",
        "图编排层与多 Agent 协作层的关系也需要区分。LangGraph 负责定义节点和条件边，相当于流程骨架；Supervisor 和各专家 Agent 负责具体业务判断，相当于流程中的执行者。把这两层分开后，后续如果更换某个 Agent 的实现，不需要重写整张图；如果调整流程路线，也不一定要修改工具本身。",
        "工具与知识层是系统能力最集中的部分。Qwen 提供开放视觉理解，PatchCore 提供本地热力图，专业 HTTP 服务提供外部检测扩展，RAG 提供相似案例和知识上下文，MMAD 分析提供结构化任务框架。这些能力都通过统一的 ToolResponse、DetectionResult 或上下文对象进入流程，避免每个模块输出完全不同格式。",
    ],
    "主运行流程设计": [
        "在首轮检测中，系统通常经历“接收任务、加载数据、规划检测、执行视觉分析、合并结果、反思、回答”几个阶段。若用户上传了图像，VisionAgent 是最常见的执行节点；若用户只输入文本问题，则系统可能跳过视觉检测或进入知识问答。这样的流程设计让同一个后端既能处理图像检测，也能支持检测后的多轮问答。",
        "在追问流程中，系统的关键不是重新上传图片，而是恢复上一轮任务。用户输入新的 question 后，run_chat 或 stream_continue_detection 会把问题写入现有状态，Supervisor 再根据新问题判断是否需要 KnowledgeAgent、ReportAgent 或重新调用 VisionAgent。也就是说，追问不是简单把旧 answer 再发一遍，而是基于已有上下文进行新一轮规划。",
        "在报告流程中，系统会检查 report_requested 标记。若该标记为真，answer 节点之后进入 report 节点；否则流程结束。这个设计避免了每次检测都生成长报告，也让前端可以通过按钮显式触发报告。对于用户来说，这比默认输出一大段报告更灵活；对于系统来说，也减少了不必要的模型调用和上下文组织开销。",
    ],
    "API接口与任务模型实现": [
        "DetectionTask 的设计体现了系统对输入的统一抽象。无论用户通过 /v1/detect、/v1/stream 还是后续追问进入系统，最终都需要落到 task_id、asset_id、question、input_type 和 parameters 等字段上。task_id 负责区分任务，asset_id 用于关联被检测对象，question 保存用户意图，parameters 则承载图像、后端选择、类别、阈值和 few-shot 上下文等扩展信息。",
        "DetectionResult 的设计则体现了输出的统一抽象。status 表示执行成功或失败，anomalies 保存异常候选，summary 和 answer 面向用户展示，metadata 保存后端和可视化信息。这样，前端可以根据 metadata 判断是否显示热力图，KnowledgeAgent 可以根据 anomalies 构造查询，ReportAgent 可以根据 result 组织报告。统一模型让不同模块之间的协作更稳定。",
        "RAG 相关 schema 单独列出，是因为它的请求和响应与检测结果不同。RagBuildStartResponse 强调任务启动，RagBuildStatusResponse 强调进度，RagQueryResponse 强调检索结果和 prompt_context。把这些 schema 写清楚，有助于论文第四章说明接口不是零散堆叠，而是围绕不同任务生命周期设计的。",
    ],
    "图编排与Supervisor实现": [
        "在代码实现中，图编排节点本身尽量保持轻量，复杂业务逻辑交给 Agent 或工具处理。例如 load_data 只检查任务是否存在并写入日志；真正决定执行哪个专家的是 supervisor_plan；真正生成用户回答的是 answer_node。这样的职责划分让每个节点更容易测试，也让出错时更容易判断问题发生在哪一层。",
        "Supervisor 的执行计划可以理解为系统的短期任务安排。计划中通常包含 step id、agent 名称、任务描述和依赖信息。执行过程中，step_status、step_attempts 和 step_outputs 会记录每一步状态。前端时间线之所以能够展示多 Agent 执行状态，正是因为这些执行事件被保存下来，而不是只在日志中一闪而过。",
        "self_reflect 的意义在于给系统一次自检机会。若视觉结果为空但用户要求定位，系统可以判断是否需要补充提示；若某个 step 失败，系统可以决定是否重试；若需要人工确认，则进入 wait_user。没有这个节点，流程就会更像固定脚本，遇到不确定结果时只能直接结束。",
    ],
    "视觉检测与PatchCore实现": [
        "Qwen 工具的 prompt 中包含了 normal_decision_policy，这一点很重要。工业图像中很多正常结构看起来也可能具有边缘、纹理和高光，如果提示词没有明确说明“不要强行寻找异常”，模型就容易为了满足用户问题而输出疑似缺陷。系统通过正常样本策略、分数阈值和文本过滤三层控制，尽量让通用视觉模型在工业检测中更稳健。",
        "PatchCore 的实现细节可以分为特征提取、距离计算、热力图生成和区域解释。特征提取使用 ResNet 等 backbone 的中间层，距离计算使用测试 patch 与 memory bank 的最近距离，热力图通过插值恢复到原图大小，区域解释则基于连通域提取 bbox。论文中写这些步骤时不需要展开复杂公式，但要把每一步输入输出说明清楚。",
        "热力图输出路径也是系统联动的关键。后端把 heatmap、mask 和 overlay 保存到 data/heatmaps 下，并通过 FastAPI 静态挂载暴露给前端。前端拿到路径后即可渲染预览。这样，PatchCore 的算法输出就从本地文件变成了用户可见的检测证据。",
    ],
    "RAG知识库与Few-shot实现": [
        "RAG 模块的核心不是“搜到几条文本”，而是把检索结果变成模型可用的上下文。Retriever 返回 rows 后，KnowledgeAgent 会把这些 rows 格式化为 prompt_context；FewShotPromptBuilder 则把正常和异常案例整理成带 category、type、severity、similarity 和 description 的文本块。这样的上下文比原始路径或 JSON 更适合注入模型。",
        "在线反馈入库功能为后续闭环学习预留了入口。用户可以把某次检测结果、图像路径、类别、描述、置信度和是否异常写入知识库。当置信度达到阈值时，系统接受样本并写入向量库；置信度过低则跳过。这一机制目前仍是轻量实现，但它说明系统设计已经考虑到知识库不是一次性构建，而是可以随着使用逐步增长。",
        "在论文中，RAG 的作用需要写得准确。它不能保证视觉模型一定判断正确，也不能代替人工标注；它的作用是提供相似案例、正常/异常参考和知识解释。把 RAG 定位为辅助证据，而不是最终裁判，既符合实际代码，也更符合工业应用中的谨慎原则。",
    ],
    "状态记忆与追问续接实现": [
        "checkpoint 后端的选择影响系统恢复能力。开发阶段使用 memory 后端简单方便，但服务重启后状态会丢失；需要持久化时可以使用 postgres 后端。论文中可以说明系统支持两类 checkpoint 配置，但不要写成已经完成复杂数据库集群部署。这样既体现了扩展设计，又不会夸大运行环境。",
        "共享上下文 SharedContext 是 Agent 协作的重要容器。VisionAgent 写入视觉结果，KnowledgeAgent 写入知识分析，ClarificationAgent 写入 pending_question，ReportAgent 再读取这些内容组织报告。如果每个 Agent 只把结果写到自己的私有字段，最终回答就很难综合多方信息。共享上下文让不同 Agent 的输出能被统一汇总。",
        "状态压缩还关系到回答质量。长对话如果完全不压缩，模型可能被过多历史消息干扰；如果压缩过度，又可能丢失关键检测结论。因此系统采用 conversation_summary 加最近对话的方式。这种折中和用户前面提到的“上下文怎么处理、如何压缩”正好对应，可以在论文中作为状态记忆模块的重点实现细节。",
    ],
    "前端交互与报告生成实现": [
        "前端的 Markdown 渲染能力也有实际意义。大模型和报告模块常常返回带标题、列表和强调的文本，如果前端只按普通字符串展示，阅读体验会比较差。支持 Markdown 后，报告内容可以分段展示，缺陷分析和建议动作也更清晰。论文中可以把它写成结果呈现优化，而不是单纯界面美化。",
        "多 Agent 时间线适合用于长任务排障。比如 RAG 建库时间较长、PatchCore 类别未训练、模型返回格式错误时，用户可以通过时间线和日志判断任务停在哪一步。此前前端只在运行摘要中体现 Agent 状态，但没有实时渲染到时间线；修复后，时间线更适合展示多 Agent 系统的运行过程。这个改动体现了系统可观测性的提升。",
        "报告生成还可以和论文截图结合。第五章后续可以放置一张完整报告截图，展示异常结论、热力图、知识分析和建议动作如何组合。需要注意的是，截图必须来自真实运行页面，不能使用静态模拟图。这样第五章的实验材料才和系统实现保持一致。",
    ],
    "测试环境与数据来源": [
        "测试环境中还需要记录模型根目录和数据集根目录。PatchCore 训练时，dataset_root 决定从哪里读取 train/good；model_root 决定 memory_bank.pt 和 metadata.json 保存到哪里；运行服务时，APP_PATCHCORE_MODEL_ROOT 又决定后端从哪里加载模型。如果这些路径写不清楚，别人复现实验时很容易出现“类别已经训练但服务找不到模型”的问题。",
        "RAG 数据集路径同样需要明确。项目中可以通过环境变量或请求参数指定 dataset_root，不同机器上的路径可能不同。例如本地 Windows 路径和 AutoDL Linux 路径完全不一样。论文中应写“以实际运行配置为准”，并在测试环境表中列出示例路径，而不是把某一台机器路径写成唯一标准。",
    ],
    "接口与功能测试": [
        "前端测试可以围绕用户操作设计，而不是只列接口名称。第一步，打开前端页面并确认后端连通；第二步，上传图片并选择后端；第三步，观察流式输出和 Agent 时间线；第四步，查看异常列表和热力图；第五步，发起追问并检查是否引用上一轮结果；第六步，触发 RAG 建库或查询。这样的测试顺序更贴近真实使用。",
        "后端测试可以结合已有 pytest 用例。项目 tests 目录中已经包含 API、chat flow、graph runtime、memory、MMAD importer、multi-agent 等相关测试文件。论文中不需要逐个展开代码，但可以说明单元测试和接口测试覆盖了基础模型、状态恢复、图运行和多 Agent 行为，为系统修改提供回归保障。",
    ],
    "PatchCore异常定位实验": [
        "实验记录中应保留训练参数。比如 image_size 为 256 和 128 时，memory bank 的大小、速度和定位效果都可能不同；是否使用 pretrained_backbone 也会影响特征表达。前期你在服务器上使用 CPU 和预训练 ResNet18 训练 bottle 类别，得到 memory_bank_size 和 recommended_threshold，这类记录就可以整理进实验表。",
        "推理结果需要结合 good 样本和异常样本一起看。只看异常图能检测出来，并不能说明系统误报率低；只看 good 图不误报，也不能说明系统能发现缺陷。论文中至少应选一组 good 和一组 defective 样本对比，展示阈值变化对结果的影响，再配合热力图截图说明异常位置是否合理。",
    ],
    "RAG检索与Few-shot实验": [
        "RAG 建库实验可以记录阶段状态，例如 pending、running、success 或 error。异步接口返回 task_id 后，前端轮询 status 接口获取 percent、processed、total 和 message。这个过程可以作为一张运行截图或表格，证明 RAG 不是一次性脚本，而是被接入前端的可操作功能。",
        "Few-shot 实验可以检查 metadata 是否记录筛选结果。理想情况下，检测或知识分析结果中应能看到 few_shot_examples、few_shot_context、normal 样本和 anomaly 样本。只要这些信息进入 metadata，就说明“RAG 相似案例 -> 样本筛选 -> prompt 构造 -> Agent 注入”这个闭环已经在系统里留下可追踪证据。",
    ],
    "多Agent流程与追问测试": [
        "多 Agent 流程测试还应覆盖失败场景。例如，当 PatchCore 类别未训练时，系统应提示缺少 memory bank，而不是进入无限重试；当 RAG 检索为空时，系统应继续完成基础检测，而不是让整个任务失败；当视觉模型返回无法解析的 JSON 时，系统应记录错误并给出明确提示。失败场景比成功场景更能体现系统鲁棒性。",
        "追问测试的评价标准可以分为三点。第一，回答是否引用上一轮图像和异常结果；第二，是否根据新问题选择合适 Agent；第三，是否避免直接复制上一轮答案。如果这三点成立，说明追问链路真正经过了状态恢复和重新规划。否则，即使前端看起来能发消息，也不能说明多轮上下文处理是完整的。",
    ],
}


def paragraph_text(paragraph) -> str:
    return paragraph.text.strip()


def style_name(paragraph) -> str:
    try:
        return paragraph.style.name.lower()
    except Exception:
        return ""


def find_body_heading(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph_text(paragraph) == text and style_name(paragraph) in {H1_STYLE, H2_STYLE}:
            return paragraph
    raise RuntimeError(f"heading not found: {text}")


def find_next_body_heading(doc: Document, heading_para):
    paragraphs = doc.paragraphs
    start = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph._element is heading_para._element
    )
    for paragraph in paragraphs[start + 1 :]:
        if style_name(paragraph) in {H1_STYLE, H2_STYLE} or paragraph_text(paragraph) == "参考文献":
            return paragraph
    raise RuntimeError(f"next heading not found after: {paragraph_text(heading_para)}")


def add_para_before(ref_para, text: str):
    para = ref_para.insert_paragraph_before(text)
    try:
        para.style = BODY_STYLE
    except Exception:
        pass
    for run in para.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: deepen_thesis_working_draft.py <docx_path>")
        return 2
    path = Path(sys.argv[1])
    doc = Document(str(path))
    for heading, paragraphs in DEEPEN_BLOCKS.items():
        anchor = find_body_heading(doc, heading)
        stop = find_next_body_heading(doc, anchor)
        for text in paragraphs:
            add_para_before(stop, text)
    doc.save(str(path))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
