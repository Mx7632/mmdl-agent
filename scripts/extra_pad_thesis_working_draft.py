from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


BODY_STYLE = "论文正文"
H1_STYLE = "heading 1"
H2_STYLE = "heading 2"


EXTRA: dict[str, list[str]] = {
    "本文主要工作": [
        "这些工作之间不是孤立关系。视觉检测提供第一手图像证据，RAG 提供相似案例和专业解释，状态记忆保证多轮追问能够接上前文，前端时间线让执行过程可见，报告生成则把一次检测沉淀为可复核文本。本文的写作重点也围绕这条链路展开：先说明为什么需要这种组合式系统，再说明系统如何分层、如何编排、如何落到代码，最后通过测试说明链路能够运行。",
    ],
    "RAG与少样本上下文": [
        "在工业异常检测中，正常样本的意义常常被低估。很多时候，判断一张图是否异常，关键并不是知道所有异常长什么样，而是知道同类正常产品在当前拍摄条件下应当是什么样。Few-shot 正常参考样本正是为了解决这个问题。它让模型在分析当前图像时有一个正常外观对照，从而减少把规则纹理、边缘或光照变化误判为异常的概率。",
    ],
    "功能需求分析": [
        "系统还需要支持配置化后端选择。对于没有训练 PatchCore 的类别，用户可以先使用 Qwen 后端做开放检测；对于已经训练过的 MVTec 类别，用户可以切换到 PatchCore 查看热力图；如果后续接入外部专业模型，也可以通过 HTTP 服务统一调用。这个需求对应 image_anomaly_detection.py 中的后端解析逻辑，也对应前端的后端选择控件。",
    ],
    "系统总体架构设计": [
        "从部署角度看，当前系统可以在本地运行，也可以把后端部署到 AutoDL 服务器，再通过 SSH 端口转发让本地前端访问远程 API。这个过程虽然不属于论文系统的核心功能，但说明 API 层和前端层相对解耦。前端只要能访问后端地址，就可以调用检测、RAG 和 PatchCore 类别接口，不必和模型训练环境部署在同一台机器上。",
    ],
    "主运行流程设计": [
        "主流程中还存在一个容易被忽略的设计点：最终回答不是由每个 Agent 自己直接返回给用户，而是由 AnswerNode 统一生成。这样做可以集中处理语言风格、失败状态、RAG 上下文、记忆摘要和 MMAD 结构。如果让每个 Agent 各自返回最终文本，系统很容易出现回答风格不一致、失败信息被覆盖或上下文遗漏的问题。",
    ],
    "视觉检测与PatchCore实现": [
        "PatchCore 的训练过程也提醒我们，工业异常检测并不一定需要大量异常样本。对于很多产品，异常样本本身稀缺且类型未知，而正常样本更容易收集。利用正常样本建立特征库，再检测偏离正常特征的区域，符合工业场景的数据分布。本文选择先在 MVTec 上训练和验证，也是因为该数据集提供了清晰的正常样本目录和异常测试目录。",
    ],
    "RAG知识库与Few-shot实现": [
        "RAG 查询结果进入 few-shot 后，还会在 metadata 中留下筛选结果。这个细节对调试很有帮助：如果模型判断异常，开发者可以回头检查它参考了哪些正常样本和异常样本；如果结果不合理，也能判断是检索结果偏了，还是模型没有正确利用上下文。这样的可追踪性比只把检索结果塞进 prompt 更重要。",
    ],
    "状态记忆与追问续接实现": [
        "状态记忆还承担了防止重复回答的作用。用户追问时，如果系统只是读取上一轮 answer 并原样返回，就不能称为真正的多轮推理。正确做法是将新问题写入 conversation_history，让 Supervisor 重新评估任务意图，再决定是否调用 KnowledgeAgent、ReportAgent 或重新执行视觉分析。这个流程让追问具备新的计算过程，而不是简单复述。",
    ],
    "前端交互与报告生成实现": [
        "前端的 RAG 页面同样服务论文实验。通过页面触发建库和查询，用户可以直接截图保存建库进度、查询结果和返回的相似样本。相比命令行输出，前端截图更适合放入第五章作为系统运行证据。后续如果需要补图，可以优先补充 RAG 建库成功页、查询结果页和 few-shot 样本展示页。",
    ],
    "接口与功能测试": [
        "功能测试还应关注数据流是否完整。例如，上传图像后，image_base64 是否进入 DetectionTask；VisionAgent 是否把 anomalies 写入 DetectionResult；KnowledgeAgent 是否能读取 anomalies 构造查询；AnswerNode 是否能综合 result 和 knowledge context；前端是否能展示 metadata 中的热力图路径。只有这些数据流都打通，系统才算真正形成闭环。",
    ],
    "PatchCore异常定位实验": [
        "在撰写实验结果时，可以把 PatchCore 分为“训练成功”“推理成功”“定位合理”“阈值仍需标定”四个层次。训练成功说明 memory bank 能生成；推理成功说明工具链可调用；定位合理需要结合热力图截图人工观察；阈值标定则说明系统仍有改进空间。这样的分析比简单写“检测成功”更细，也更容易体现工程判断。",
    ],
    "RAG检索与Few-shot实验": [
        "Few-shot 实验还可以观察正常样本和异常样本比例。当前选择器默认取少量正常样本和异常样本，避免 prompt 过长。如果全部塞入检索结果，模型上下文会变得冗余，重要样本反而被淹没。选择少而有代表性的案例，是 few-shot prompt 构造的基本原则，也是本文系统选择器存在的原因。",
    ],
    "多Agent流程与追问测试": [
        "对于多 Agent 系统，测试不只看最终答案，还要看中间过程是否符合预期。比如用户问“有没有使用 RAG”，系统应当能够从上下文或 metadata 中说明是否检索过知识库；用户问“为什么判断为异常”，系统应当调用或复用 KnowledgeAgent 的分析结果；用户要求“重新执行”，系统应当进入新一轮规划而不是仅返回旧结果。这些都是多 Agent 架构有效性的体现。",
    ],
    "后续优化方向": [
        "最后，论文后续工作还可以补充用户反馈闭环和更细粒度权限控制。当前系统更偏单用户实验环境，后续如果面向真实工厂，需要考虑不同人员对模型训练、知识库更新、报告审核和历史记录的操作边界。不过这些属于工程扩展内容，本文只将其作为展望，不写成已经完成的系统功能。",
    ],
}


def paragraph_text(paragraph) -> str:
    return paragraph.text.strip()


def style_name(paragraph) -> str:
    try:
        return paragraph.style.name.lower()
    except Exception:
        return ""


def find_body_heading(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph_text(paragraph) == text and style_name(paragraph) in {H1_STYLE, H2_STYLE}:
            return paragraph
    raise RuntimeError(f"heading not found: {text}")


def find_next_body_heading(doc: Document, heading_para):
    paragraphs = doc.paragraphs
    start = next(index for index, paragraph in enumerate(paragraphs) if paragraph._element is heading_para._element)
    for paragraph in paragraphs[start + 1 :]:
        if style_name(paragraph) in {H1_STYLE, H2_STYLE} or paragraph_text(paragraph) == "参考文献":
            return paragraph
    raise RuntimeError("next heading not found")


def add_para_before(ref_para, text: str):
    para = ref_para.insert_paragraph_before(text)
    try:
        para.style = BODY_STYLE
    except Exception:
        pass
    for run in para.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: extra_pad_thesis_working_draft.py <docx_path>")
        return 2
    path = Path(sys.argv[1])
    doc = Document(str(path))
    for heading, paragraphs in EXTRA.items():
        anchor = find_body_heading(doc, heading)
        stop = find_next_body_heading(doc, anchor)
        for text in paragraphs:
            add_para_before(stop, text)
    doc.save(str(path))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
